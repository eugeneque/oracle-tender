"""Учёт расхода токенов, окно контекста моделей и карточка производителя под пачку
требований (14.10.2026)."""

from __future__ import annotations

import uuid
from datetime import datetime, timezone
from types import SimpleNamespace

import httpx
import pydantic
import pytest
from sqlalchemy import select

from app.models.ai_usage import AiUsageRecord
from app.models.job import BackgroundJob
from app.models.source import Source
from app.models.tender import Tender
from app.services import ai_client, ai_provider_service, ai_usage_service, routerai_client
from app.services.ai_context import running_job
from app.services.compliance_service import focus_context


class Answer(pydantic.BaseModel):
    text: str


def _ok_response(text: str = '{"text": "ok"}', **usage) -> httpx.Response:
    return httpx.Response(
        200,
        json={
            "model": "deepseek/deepseek-v4-pro-0813",
            "choices": [{"message": {"content": text}, "finish_reason": "stop"}],
            "usage": usage,
        },
    )


@pytest.fixture()
def routerai(monkeypatch):
    """RouterAI без сети: ответы шлюза — из списка, тела запросов сохраняются."""

    state = {"responses": [], "bodies": []}

    def fake_post(url, *, headers, json, timeout):
        state["bodies"].append(json)
        response = state["responses"].pop(0)
        if isinstance(response, Exception):
            raise response
        return response

    monkeypatch.setattr(routerai_client.httpx, "post", fake_post)
    monkeypatch.setattr(
        routerai_client,
        "get_routerai_credentials",
        lambda db, provider="deepseek": ("k", "deepseek/deepseek-v4-pro-0813", "https://x"),
    )
    monkeypatch.setattr(routerai_client.time, "sleep", lambda seconds: None)
    return state


def _job_with_tender(db_session) -> tuple[BackgroundJob, Tender]:
    suffix = uuid.uuid4().hex[:8]
    source = Source(key=f"use_{suffix}", name="Источник", url=f"https://example.test/{suffix}", type="eis")
    db_session.add(source)
    db_session.flush()
    tender = Tender(source_id=source.id, external_id=f"U-{suffix}", title="Поставка счётчиков")
    db_session.add(tender)
    db_session.flush()
    job = BackgroundJob(kind="tender_evaluation", status="running", tender_id=tender.id)
    db_session.add(job)
    db_session.commit()
    return job, tender


def _records(db_session, job_id) -> list[AiUsageRecord]:
    return list(
        db_session.scalars(
            select(AiUsageRecord).where(AiUsageRecord.job_id == job_id).order_by(AiUsageRecord.created_at)
        )
    )


def test_call_is_recorded_with_provider_counts_and_job_tender(db_session, routerai):
    """Счётчик провайдера, кэш и стоимость сохраняются как есть; закупка и вид задачи —
    из выполняемой фоновой задачи, без протаскивания через сигнатуры сервисов."""

    job, tender = _job_with_tender(db_session)
    routerai["responses"] = [
        _ok_response(
            prompt_tokens=1200,
            completion_tokens=80,
            prompt_tokens_details={"cached_tokens": 1000},
            cost=0.42,
        )
    ]

    with running_job(datetime.now(timezone.utc), job.id):
        answer = routerai_client.run_structured(
            db_session, system_prompt="s" * 100, user_text="u" * 300, response_model=Answer
        )

    assert answer.text == "ok"
    (record,) = _records(db_session, job.id)
    assert record.provider == "deepseek"
    assert record.operation == "Answer"
    assert (record.prompt_tokens, record.completion_tokens, record.cached_tokens) == (1200, 80, 1000)
    assert record.tokens_estimated is False
    assert record.prompt_chars == 400
    assert float(record.cost) == pytest.approx(0.42)
    assert record.tender_id == tender.id and record.job_kind == "tender_evaluation"
    assert record.status == "success"


def test_every_attempt_is_recorded_and_missing_counts_are_estimated(db_session, routerai):
    """Повтор оплачивается так же, как первый запрос, — записывается каждая попытка. Сбой
    до ответа счётчика не даёт: токены оцениваются по длине текста и помечаются оценкой."""

    job, _ = _job_with_tender(db_session)
    routerai["responses"] = [
        httpx.ConnectError("нет связи"),
        _ok_response("не JSON", prompt_tokens=500, completion_tokens=10),
        _ok_response(prompt_tokens=500, completion_tokens=12),
    ]

    with running_job(datetime.now(timezone.utc), job.id):
        routerai_client.run_structured(
            db_session, system_prompt="s", user_text="у" * 239, response_model=Answer
        )

    first, second, third = _records(db_session, job.id)
    assert first.status == "error" and first.tokens_estimated is True
    assert first.prompt_tokens == round(240 / ai_provider_service.CHARS_PER_TOKEN)
    assert second.status == "invalid" and second.prompt_tokens == 500
    assert third.status == "success"


def test_output_ceiling_comes_from_model_limits(db_session, admin_user, routerai):
    """Потолок ответа — из настроек модели: на окне 32 тыс. ответ в 16 тыс. не помещается."""

    ai_provider_service.update_model_limits(
        db_session, "deepseek", context_tokens=32_768, max_output_tokens=4_000, actor=admin_user
    )
    routerai["responses"] = [_ok_response(prompt_tokens=10, completion_tokens=2)]

    routerai_client.run_structured(db_session, system_prompt="s", user_text="u", response_model=Answer)

    assert routerai["bodies"][0]["max_tokens"] == 4_000


def test_qwen_goes_through_routerai_without_reasoning(db_session, routerai):
    """Qwen3.8 27B (20.10.2026) — через тот же шлюз, с «мышлением», выключенным как у
    DeepSeek, и с окном 262 тыс. по умолчанию; запись учёта — под провайдером qwen."""

    job, _tender = _job_with_tender(db_session)
    routerai["responses"] = [_ok_response(prompt_tokens=10, completion_tokens=2)]
    with running_job(datetime.now(timezone.utc), job.id):
        routerai_client.run_structured(
            db_session, system_prompt="s", user_text="u", response_model=Answer, provider="qwen"
        )

    assert routerai["bodies"][0]["reasoning"] == {"enabled": False}
    assert ai_provider_service.get_model_limits(db_session, "qwen").context_tokens == 262_144
    (record,) = _records(db_session, job.id)
    assert record.provider == "qwen"


def test_model_limits_validation_budget_and_reset(db_session, admin_user):
    with pytest.raises(ValueError, match="половины окна"):
        ai_provider_service.update_model_limits(
            db_session, "yandex", context_tokens=16_384, max_output_tokens=9_000, actor=admin_user
        )
    with pytest.raises(ValueError, match="Окно контекста"):
        ai_provider_service.update_model_limits(
            db_session, "yandex", context_tokens=1_000, max_output_tokens=None, actor=admin_user
        )

    ai_provider_service.update_model_limits(
        db_session, "yandex", context_tokens=16_384, max_output_tokens=4_000, actor=admin_user
    )
    expected = int((16_384 - 4_000 - ai_provider_service.SAFETY_TOKENS) * ai_provider_service.CHARS_PER_TOKEN)
    assert ai_provider_service.input_budget_chars(db_session, "yandex") == expected

    rows = ai_provider_service.update_model_limits(
        db_session, "yandex", context_tokens=None, max_output_tokens=None, actor=admin_user
    )
    yandex = next(row for row in rows if row["provider"] == "yandex")
    assert yandex["is_default"] is True and yandex["context_tokens"] == 32_768


def test_document_chunk_shrinks_to_fit_a_small_window(db_session, monkeypatch):
    """Кусок документации — не длиннее, чем помещается в окно активной модели."""

    monkeypatch.setattr(ai_client, "input_budget_chars", lambda db: 9_000)
    assert ai_client.fit_chars(db_session, 14_000, reserved=3_000) == 6_000
    monkeypatch.setattr(ai_client, "input_budget_chars", lambda db: 80_000)
    assert ai_client.fit_chars(db_session, 14_000, reserved=3_000) == 14_000


def _requirement(text: str):
    return SimpleNamespace(normalized_text=text, text=text, parameter_no=None)


def test_short_card_is_sent_as_is():
    context = "[si_type] номер в Госреестре СИ 61891-15\n[catalog] М1: Связь — Модем: GSM"
    assert focus_context(context, [_requirement("Модем GSM")], 10_000) == context


def test_long_card_keeps_relevant_and_pinned_facts_within_limit():
    """Пачке о GSM-модеме нужна строка про модем, а не сотни строк о прочем; факты [tz] и
    [registry] остаются всегда — по ним правила промпта решают «нет — не соответствует»."""

    lines = ["[tz] ТЗ составлено по модели М1 этого производителя"]
    lines += [f"[catalog] М1: Общие — Параметр {index}: значение {index} при температуре" for index in range(400)]
    lines.append("[catalog] М1: Связь — Модем: GSM/GPRS, 2G")
    lines.append("[registry] М1: реестр ПП 719 — запись действует")
    context = "\n".join(lines)

    focused = focus_context(context, [_requirement("Наличие встроенного GSM-модема")], 3_000)

    assert len(focused) <= 3_000
    assert "Модем: GSM/GPRS" in focused
    assert "[tz] ТЗ составлено" in focused and "[registry] М1" in focused


def test_long_block_is_cut_to_pieces_with_its_header():
    """Из длинной выдержки «Описания типа» берутся только нужные куски — с заголовком
    блока и пометкой пропуска, чтобы модель видела, что выдержка неполная."""

    filler = [f"Пункт {index}: описание конструкции корпуса и пломбировки" for index in range(120)]
    block = "\n".join(
        ["[si_type] Выдержка из «Описания типа» 61891-15:", *filler[:60], "G — GSM-модуль связи", *filler[60:]]
    )
    other = "\n".join(f"[catalog] М1: Прочее — Поле {index}: {index}" for index in range(200))
    focused = focus_context(f"{block}\n{other}", [_requirement("GSM-модуль")], 2_500)

    assert "[si_type] Выдержка из «Описания типа» 61891-15:" in focused
    assert "G — GSM-модуль связи" in focused
    assert "[…]" in focused


def test_summary_groups_steps_and_averages_per_tender(db_session):
    """Четыре схемы AI-оценки — один шаг разбора; среднее на закупку — по видам задач."""

    provider = f"t{uuid.uuid4().hex[:8]}"
    tender_a, tender_b = uuid.uuid4(), uuid.uuid4()
    for operation, tender_id, prompt in (
        ("TaskAnswer", tender_a, 1000),
        ("ResumeAnswer", tender_a, 500),
        ("ComplianceResult", tender_b, 3000),
    ):
        db_session.add(
            AiUsageRecord(
                provider=provider,
                model="m",
                operation=operation,
                job_kind="tender_full_review",
                tender_id=tender_id,
                prompt_tokens=prompt,
                completion_tokens=100,
                status="success",
            )
        )
    db_session.commit()

    today = datetime.now(ai_usage_service.MOSCOW_TZ).date()
    summary = ai_usage_service.usage_summary(db_session, date_from=today, date_to=today, provider=provider)

    assert summary["totals"]["calls"] == 3
    assert summary["totals"]["total_tokens"] == 4800
    assert summary["totals"]["tenders"] == 2
    steps = {row["label"]: row for row in summary["by_operation"]}
    assert steps["AI-оценка по профилю"]["calls"] == 2
    assert steps["Матрица соответствия"]["max_prompt_tokens"] == 3000
    (kind,) = summary["by_job_kind"]
    assert kind["label"] == "Полный разбор закупки" and kind["avg_tokens_per_tender"] == 2400


def test_usage_endpoints_are_admin_only_and_check_period(client, admin_token):
    headers = {"Authorization": f"Bearer {admin_token}"}

    summary = client.get("/ai-usage/summary", headers=headers)
    assert summary.status_code == 200, summary.text
    assert {"totals", "by_day", "by_model", "by_operation"} <= summary.json().keys()

    wrong = client.get(
        "/ai-usage/summary", params={"date_from": "2026-10-10", "date_to": "2026-10-01"}, headers=headers
    )
    assert wrong.status_code == 400 and "позже" in wrong.json()["detail"]

    assert client.get("/ai-usage/calls", headers=headers).status_code == 200
    assert client.get("/ai-usage/summary").status_code == 401

    limits = client.get("/integrations/ai-limits", headers=headers)
    assert {row["provider"] for row in limits.json()} == {"yandex", "deepseek", "qwen"}


def test_user_usage_lists_steps_and_tenders_with_links(db_session, admin_user, client, admin_token):
    """«На что тратит токены пользователь»: шаги разбора и разобранные им закупки, сначала
    последние; чужие обращения и обращения расписания в его разбивку не попадают."""

    provider = f"t{uuid.uuid4().hex[:8]}"
    _, tender = _job_with_tender(db_session)
    deleted_tender = uuid.uuid4()
    for operation, tender_id, user_id in (
        ("RequirementsResult", tender.id, admin_user.id),
        ("ComplianceResult", tender.id, admin_user.id),
        ("TaskAnswer", deleted_tender, admin_user.id),
        ("RelevanceAnswer", None, None),
    ):
        db_session.add(
            AiUsageRecord(
                provider=provider,
                model="m",
                operation=operation,
                job_kind="tender_full_review" if tender_id else None,
                tender_id=tender_id,
                user_id=user_id,
                prompt_tokens=1000,
                completion_tokens=100,
                status="success",
            )
        )
    db_session.commit()

    today = datetime.now(ai_usage_service.MOSCOW_TZ).date()
    usage = ai_usage_service.user_usage(
        db_session, user_key=str(admin_user.id), date_from=today, date_to=today, provider=provider
    )
    assert usage["totals"]["calls"] == 3
    assert {row["label"] for row in usage["by_operation"]} == {
        "Извлечение требований",
        "Матрица соответствия",
        "AI-оценка по профилю",
    }
    assert usage["tenders_total"] == 2
    by_id = {row["tender_id"]: row for row in usage["tenders"]}
    assert by_id[str(tender.id)]["external_id"] == tender.external_id
    assert by_id[str(tender.id)]["operations"] == ["Извлечение требований", "Матрица соответствия"]
    assert by_id[str(deleted_tender)]["exists"] is False

    system = ai_usage_service.user_usage(
        db_session, user_key="system", date_from=today, date_to=today, provider=provider
    )
    assert system["totals"]["calls"] == 1 and system["tenders"] == []

    headers = {"Authorization": f"Bearer {admin_token}"}
    assert client.get(f"/ai-usage/users/{admin_user.id}", headers=headers).status_code == 200
    bad = client.get("/ai-usage/users/not-a-uuid", headers=headers)
    assert bad.status_code == 400 and "пользователь" in bad.json()["detail"]



def test_usage_xlsx_has_sheet_per_section_with_speed_and_cache(db_session, client, admin_token):
    """Выгрузка: лист на раздел; скорость и время ответа — по успешным обращениям."""

    from io import BytesIO

    from openpyxl import load_workbook

    provider = f"t{uuid.uuid4().hex[:8]}"
    for status, completion, duration in (("success", 500, 2000), ("error", 0, 100)):
        db_session.add(
            AiUsageRecord(
                provider=provider,
                model="m",
                operation="ComplianceResult",
                prompt_tokens=2000,
                completion_tokens=completion,
                cached_tokens=500,
                duration_ms=duration,
                status=status,
            )
        )
    db_session.commit()

    today = datetime.now(ai_usage_service.MOSCOW_TZ).date()
    content, file_name, rows = ai_usage_service.build_usage_xlsx(
        db_session, date_from=today, date_to=today, provider=provider
    )
    assert file_name.endswith(".xlsx") and rows == 2
    book = load_workbook(BytesIO(content))
    assert book.sheetnames == [
        "Итог", "По дням", "Модели", "Шаги разбора", "Виды разбора", "Пользователи", "Закупки", "Обращения",
    ]

    steps = book["Шаги разбора"]
    header = [cell.value for cell in steps[1]]
    step = dict(zip(header, [cell.value for cell in steps[2]]))
    assert step["Шаг"] == "Матрица соответствия" and step["Обращений"] == 2 and step["Сбоев"] == 1
    # 1000 из 4000 входящих — из кэша; сбой не тянет вниз время ответа и скорость.
    assert step["Доля кэша"] == 0.25
    assert step["Ср. время ответа, с"] == 2.0 and step["Скорость вывода, ток/с"] == 250.0
    assert steps.freeze_panes == "B2" and steps.auto_filter.ref

    calls = book["Обращения"]
    assert calls.max_row == 3 and isinstance(calls["A2"].value, datetime)

    headers = {"Authorization": f"Bearer {admin_token}"}
    response = client.get("/ai-usage/export.xlsx", headers=headers)
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "sova_ai_usage_" in response.headers["content-disposition"]
    assert client.get("/ai-usage/export.xlsx").status_code == 401
