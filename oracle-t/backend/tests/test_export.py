"""Выгрузка тендеров в Excel (раздел 5.7 ТЗ, Приложение D)."""

import io
import uuid
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal

import openpyxl
from sqlalchemy import select

from app.models.log import Log
from app.models.region import RegionResponsible
from app.models.source import Source
from app.models.tender import Tender
from app.services.export_service import (
    COLUMNS,
    FIRST_DATA_ROW,
    HEADER_ROW,
    build_tenders_workbook,
    export_tenders,
)
from app.services.tender_service import TenderFilters

# Заголовки Приложения D в точном порядке — тест сверяет их дословно, потому что заказчик
# принимает выгрузку по позициям приложения, а не по смыслу названий.
APPENDIX_D_HEADERS = [
    "Статус",
    "Начало приёма заявок",
    "Окончание приёма заявок",
    "Осталось дней",
    "Месяц",
    "Сумма",
    "Валюта",
    "Тип конкурса",
    "Способ закупки",
    "Организатор",
    "ФО",
    "Регион организатора",
    "Регион поставки",
    "Ответственный за регион",
    "Руководитель ответственный за регион",
    "Наименование ЭТП",
    "Ссылка",
    "ID сделки",
    "Комментарий",
]


def _auth_headers(token: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {token}"}


def _make_source(db) -> Source:
    source = Source(
        key=f"export_{uuid.uuid4().hex[:8]}",
        name="Площадка для выгрузки",
        url="https://example.test",
        type="etp_federal_commercial",
    )
    db.add(source)
    db.commit()
    db.refresh(source)
    return source


def _make_tender(db, source, **kwargs) -> Tender:
    tender = Tender(
        source_id=source.id,
        external_id=kwargs.pop("external_id", f"EXT-{uuid.uuid4().hex[:8]}"),
        title=kwargs.pop("title", "Поставка приборов учёта"),
        currency=kwargs.pop("currency", "RUB"),
        **kwargs,
    )
    db.add(tender)
    db.commit()
    db.refresh(tender)
    return tender


def _sheet(content: bytes):
    return openpyxl.load_workbook(io.BytesIO(content)).active


def test_export_requires_auth(client):
    assert client.get("/export/tenders.xlsx").status_code == 401


def test_export_contains_all_19_appendix_d_fields(db_session):
    source = _make_source(db_session)
    _make_tender(db_session, source)

    content, rows = build_tenders_workbook(
        db_session, TenderFilters(source_keys=[source.key])
    )
    sheet = _sheet(content)

    headers = [sheet.cell(HEADER_ROW, column).value for column in range(1, len(COLUMNS) + 1)]
    assert headers == APPENDIX_D_HEADERS
    assert rows == 1


def test_export_header_carries_generation_date_and_period(db_session):
    """Дата формирования и период отбора в шапке файла (REQ-5.7-09, 5.7-10)."""

    source = _make_source(db_session)
    _make_tender(db_session, source, publish_date=date(2026, 7, 15))

    content, _ = build_tenders_workbook(
        db_session,
        TenderFilters(
            source_keys=[source.key],
            publish_date_from=date(2026, 7, 1),
            publish_date_to=date(2026, 7, 31),
        ),
    )
    sheet = _sheet(content)

    assert sheet.cell(2, 1).value == "Дата формирования:"
    assert datetime.now().strftime("%d.%m.%Y") in sheet.cell(2, 2).value
    assert sheet.cell(3, 2).value == "дата размещения с 01.07.2026 по 31.07.2026"


def test_export_without_date_filter_says_so_explicitly(db_session):
    """Без ограничения по датам период называется прямо — иначе выгрузку за всю базу
    легко принять за отчёт за месяц."""

    source = _make_source(db_session)
    _make_tender(db_session, source)

    content, _ = build_tenders_workbook(db_session, TenderFilters(source_keys=[source.key]))

    assert "весь период" in _sheet(content).cell(3, 2).value


def test_export_freezes_header_and_sets_autofilter(db_session):
    """Закреплённая шапка и автофильтр (REQ-5.7-04, 5.7-05)."""

    source = _make_source(db_session)
    _make_tender(db_session, source)

    content, _ = build_tenders_workbook(db_session, TenderFilters(source_keys=[source.key]))
    sheet = _sheet(content)

    assert sheet.freeze_panes == f"A{FIRST_DATA_ROW}"
    assert sheet.auto_filter.ref is not None
    assert sheet.auto_filter.ref.startswith(f"A{HEADER_ROW}:")


def test_export_applies_conditional_formatting(db_session):
    """Условное форматирование по сроку подачи и проценту победителя (REQ-5.7-06, 5.7-07).

    Правила именно условные, а не заливка при генерации: пользователь пересортировывает
    выгрузку руками, и цвет должен следовать за значением, а не оставаться на строке.
    """

    source = _make_source(db_session)
    _make_tender(
        db_session, source, application_end=datetime.now(timezone.utc) + timedelta(days=2)
    )

    content, _ = build_tenders_workbook(db_session, TenderFilters(source_keys=[source.key]))
    sheet = _sheet(content)

    ranges = {str(rng.sqref).split(":")[0][0] for rng in sheet.conditional_formatting}
    # D — «Осталось дней», T — колонка процента победителя, добавленная после 19 полей.
    assert "D" in ranges
    assert "T" in ranges


def test_export_row_values_follow_appendix_d(db_session):
    source = _make_source(db_session)
    deadline = datetime.now(timezone.utc) + timedelta(days=3)
    _make_tender(
        db_session,
        source,
        external_id="EXPORT-42",
        status="collecting_bids",
        tender_type="supply_only",
        price=Decimal("1234.56"),
        publish_date=date(2026, 4, 10),
        application_end=deadline,
        organizer_name="ГКУ «Заказчик»",
        procurement_method="Электронный аукцион",
        source_url="https://example.test/purchase/42",
        ai_comment="Закупка счётчиков, подходит по классу точности",
    )

    content, _ = build_tenders_workbook(db_session, TenderFilters(source_keys=[source.key]))
    sheet = _sheet(content)
    row = FIRST_DATA_ROW

    assert sheet.cell(row, 1).value == "Сбор заявок"
    assert sheet.cell(row, 4).value == 3  # осталось дней
    assert sheet.cell(row, 5).value == "апрель"
    assert sheet.cell(row, 6).value == 1234.56
    assert sheet.cell(row, 7).value == "RUB"
    assert sheet.cell(row, 8).value == "Поставка ИПУ"
    assert sheet.cell(row, 9).value == "Электронный аукцион"
    assert sheet.cell(row, 10).value == "ГКУ «Заказчик»"
    assert sheet.cell(row, 16).value == source.name
    assert sheet.cell(row, 18).value == "EXPORT-42"
    assert sheet.cell(row, 19).value == "Закупка счётчиков, подходит по классу точности"


def test_export_link_cell_is_a_hyperlink(db_session):
    """Ссылка отдаётся именно гиперссылкой (REQ-5.7-08), а не текстом URL."""

    source = _make_source(db_session)
    _make_tender(db_session, source, source_url="https://example.test/purchase/7")

    content, _ = build_tenders_workbook(db_session, TenderFilters(source_keys=[source.key]))
    cell = _sheet(content).cell(FIRST_DATA_ROW, 17)

    assert cell.hyperlink is not None
    assert cell.hyperlink.target == "https://example.test/purchase/7"


def test_delivery_region_falls_back_to_organizer_region(db_session):
    """Регион поставки дублирует регион организатора, если отдельно не указан
    (Приложение D, раздел 5.4 п.6 ТЗ). Подстановка делается в выгрузке, а не в БД: пустое
    поле в базе означает «не определено», и затирать это дублем нельзя."""

    source = _make_source(db_session)
    _make_tender(
        db_session, source, region_organizer_code="23", region_delivery_code=None
    )

    content, _ = build_tenders_workbook(db_session, TenderFilters(source_keys=[source.key]))
    sheet = _sheet(content)

    assert sheet.cell(FIRST_DATA_ROW, 12).value == "Краснодарский край"
    assert sheet.cell(FIRST_DATA_ROW, 13).value == "Краснодарский край"


def test_responsible_columns_filled_from_dictionary(db_session):
    """Два поля Приложения D берутся из справочника «регион → ответственный»."""

    source = _make_source(db_session)
    _make_tender(db_session, source, region_organizer_code="23")
    db_session.merge(
        RegionResponsible(
            region_code="23", responsible_name="Иванов И.И.", manager_name="Петров П.П."
        )
    )
    db_session.commit()

    content, _ = build_tenders_workbook(db_session, TenderFilters(source_keys=[source.key]))
    sheet = _sheet(content)

    assert sheet.cell(FIRST_DATA_ROW, 14).value == "Иванов И.И."
    assert sheet.cell(FIRST_DATA_ROW, 15).value == "Петров П.П."


def test_export_respects_filters(db_session):
    """Выгрузка повторяет срез списка: иначе отчёт расходился бы с тем, что пользователь
    только что отобрал на экране, и объяснить расхождение было бы нечем."""

    source = _make_source(db_session)
    _make_tender(db_session, source, price=Decimal("100.00"))
    _make_tender(db_session, source, price=Decimal("100000.00"))

    _, rows = build_tenders_workbook(
        db_session,
        TenderFilters(source_keys=[source.key], price_min=Decimal("1000")),
    )

    assert rows == 1


def test_export_is_logged(db_session, admin_user):
    """Экспорт данных фиксируется в журнале (REQ-5.9-07)."""

    source = _make_source(db_session)
    _make_tender(db_session, source)

    _, file_name, rows = export_tenders(
        db_session, TenderFilters(source_keys=[source.key]), admin_user
    )

    assert file_name.endswith(".xlsx")
    entry = db_session.scalar(
        select(Log)
        .where(Log.component == "export", Log.action == "export_tenders_xlsx")
        .order_by(Log.timestamp.desc())
    )
    assert entry is not None
    assert entry.user_id == admin_user.id
    assert f"Строк выгружено: {rows}" in entry.details


def test_export_endpoint_returns_xlsx_with_row_count(client, admin_token):
    response = client.get("/export/tenders.xlsx?limit=5", headers=_auth_headers(admin_token))

    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert ".xlsx" in response.headers["content-disposition"]
    # Фронт показывает «выгружено N строк», а тело ответа — бинарник: число едет заголовком.
    assert response.headers["x-exported-rows"].isdigit()
    # Файл должен открываться без ошибок (REQ-5.7-11) — если структура битая,
    # openpyxl не разберёт его.
    assert _sheet(response.content).title == "Тендеры"


def test_region_responsible_requires_admin(client, admin_token):
    """Справочник ответственных правит только администратор — это распределение работы
    между людьми, а не пользовательская настройка отображения."""

    created = client.post(
        "/users",
        headers=_auth_headers(admin_token),
        json={
            "username": f"resp_{uuid.uuid4().hex[:8]}",
            "password": "UserPass123!",
            "full_name": "Обычный пользователь",
            "role": "user",
        },
    )
    assert created.status_code == 201, created.text

    token = client.post(
        "/auth/login",
        json={"username": created.json()["username"], "password": "UserPass123!"},
    ).json()["access_token"]

    response = client.put(
        "/dictionaries/region-responsibles/23",
        headers=_auth_headers(token),
        json={"responsible_name": "Сидоров С.С."},
    )
    assert response.status_code == 403


def test_region_responsible_upsert_and_list(client, admin_token):
    response = client.put(
        "/dictionaries/region-responsibles/23",
        headers=_auth_headers(admin_token),
        json={"responsible_name": "Иванов И.И.", "manager_name": "Петров П.П."},
    )
    assert response.status_code == 200, response.text
    assert response.json()["region_name"] == "Краснодарский край"

    # Повторный PUT обновляет ту же запись, а не создаёт вторую.
    updated = client.put(
        "/dictionaries/region-responsibles/23",
        headers=_auth_headers(admin_token),
        json={"responsible_name": "Сидоров С.С.", "manager_name": None},
    )
    assert updated.status_code == 200
    assert updated.json()["responsible_name"] == "Сидоров С.С."
    assert updated.json()["manager_name"] is None

    listed = client.get("/dictionaries/region-responsibles", headers=_auth_headers(admin_token))
    assert listed.status_code == 200
    rows = {row["region_code"]: row for row in listed.json()}
    assert rows["23"]["responsible_name"] == "Сидоров С.С."


def test_region_responsible_rejects_unknown_region(client, admin_token):
    response = client.put(
        "/dictionaries/region-responsibles/99",
        headers=_auth_headers(admin_token),
        json={"responsible_name": "Иванов И.И."},
    )
    assert response.status_code == 404
