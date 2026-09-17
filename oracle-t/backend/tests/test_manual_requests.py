"""Ручные заявки: закупка, которую заказчик прислал напрямую (решение 15.09.2026)."""

import io
import uuid

from sqlalchemy import select

from app.core import jobs
from app.db.session import SessionLocal
from app.models.job import BackgroundJob, JobKind
from app.models.source import MANUAL_SOURCE_KEY, Source
from app.models.tender import Tender
from app.models.tender_document import TenderDocument
from app.models.tender_history import TenderHistoryEntry
from app.services import tender_service


def _auth_headers(token: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {token}"}


def _use_storage(monkeypatch, tmp_path) -> None:
    """Файлы заявки — во временный каталог, а не в рабочее хранилище. Подменяются все три
    места, где имя `get_storage_root` связано на импорте: сервис документов, сервис заявок и
    эндпоинт скачивания."""

    import app.api.endpoints.tenders as tenders_endpoints
    import app.services.document_service as document_service_module
    import app.services.manual_request_service as manual_module

    for module in (tenders_endpoints, document_service_module, manual_module):
        monkeypatch.setattr(module, "get_storage_root", lambda: tmp_path)


def _docx_bytes(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    from docx import Document

    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for row_index, row in enumerate(table):
            for col_index, value in enumerate(row):
                grid.cell(row_index, col_index).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_manual_source_is_seeded_and_never_polled(db_session):
    """Системный источник заявок есть в базе, но опрос площадок его не трогает: у него нет
    ни сайта, ни адаптера, и «адаптер не реализован» в журнале по нему было бы ложной
    ошибкой каждый день."""

    source = db_session.scalar(select(Source).where(Source.key == MANUAL_SOURCE_KEY))
    assert source is not None
    assert source.type == "manual"

    assert tender_service.poll_sources(db_session, [MANUAL_SOURCE_KEY]) == []


def test_create_manual_request_with_file(client, admin_token, tmp_path, monkeypatch):
    """Заявка с приложенным договором: тендер создан, файл разобран, анализ поставлен в
    очередь, комментарий попал в историю."""

    _use_storage(monkeypatch, tmp_path)
    # Задача только ставится, не выполняется: пул фоновых задач общий на весь прогон, и
    # после `lifespan`-остановки приложения в соседнем тесте он уже закрыт (см. test_jobs).
    monkeypatch.setattr(jobs._executor, "submit", lambda *args, **kwargs: None)

    contract = _docx_bytes(
        ["Проект договора на поставку приборов учёта"],
        table=[["Параметр", "Значение"], ["Класс точности", "1,0"]],
    )
    response = client.post(
        "/tenders/manual",
        headers=_auth_headers(admin_token),
        data={
            "title": "Поставка счётчиков для АО «Новосибирскэнергосбыт»",
            "customer_name": "АО «Новосибирскэнергосбыт»",
            "price": "1500000.50",
            "application_end": "2026-10-01",
            "comment": "Проект договора пришёл письмом, характеристики на стр. 84",
        },
        files=[("files", ("Проект договора.docx", contract, "application/octet-stream"))],
    )
    assert response.status_code == 201, response.text
    payload = response.json()

    tender = payload["tender"]
    assert tender["source"]["key"] == MANUAL_SOURCE_KEY
    assert tender["external_id"].startswith("ЗАЯВКА-")
    assert tender["stage"] == "under_review"
    assert tender["customer_name"] == "АО «Новосибирскэнергосбыт»"
    assert tender["price"] == "1500000.50"
    assert tender["application_end"].startswith("2026-10-01")

    assert len(payload["documents"]) == 1
    document = payload["documents"][0]
    assert document["file_name"] == "Проект договора.docx"
    assert document["parse_status"] == "success"
    assert document["has_text"] is True
    assert document["document_class"] == "contract"

    assert payload["job"] is not None
    assert payload["job"]["kind"] == JobKind.TENDER_ANALYSIS.value

    db = SessionLocal()
    try:
        stored = db.get(TenderDocument, uuid.UUID(document["id"]))
        assert stored is not None
        assert stored.source_url.startswith("upload://")
        assert (tmp_path / stored.storage_path).is_file()
        # Таблица прочитана строкой «параметр | значение», а не двумя ячейками в столбик.
        assert "Класс точности | 1,0" in (stored.extracted_text or "")

        comments = db.scalars(
            select(TenderHistoryEntry).where(
                TenderHistoryEntry.tender_id == uuid.UUID(tender["id"])
            )
        ).all()
        assert [entry.comment for entry in comments] == [
            "Проект договора пришёл письмом, характеристики на стр. 84"
        ]

        job = db.scalar(
            select(BackgroundJob).where(BackgroundJob.tender_id == uuid.UUID(tender["id"]))
        )
        assert job is not None
    finally:
        db.close()

    # Загруженный файл отдаётся из хранилища — «перекачивать» его неоткуда.
    download = client.get(
        f"/tenders/{tender['id']}/documents/{document['id']}/download",
        headers=_auth_headers(admin_token),
    )
    assert download.status_code == 200
    assert download.content == contract


def test_manual_request_without_files_does_not_enqueue_analysis(client, admin_token):
    """Без документов анализировать нечего, кроме названия: задача не ставится, заявка
    создаётся — файлы можно приложить позже."""

    response = client.post(
        "/tenders/manual",
        headers=_auth_headers(admin_token),
        data={"title": "Заявка без файлов"},
    )
    assert response.status_code == 201, response.text
    assert response.json()["documents"] == []
    assert response.json()["job"] is None


def test_manual_request_requires_title(client, admin_token):
    response = client.post(
        "/tenders/manual",
        headers=_auth_headers(admin_token),
        data={"title": "   "},
    )
    assert response.status_code == 422
    assert "наименование" in response.json()["detail"].lower()


def test_upload_documents_to_existing_tender(client, admin_token, tmp_path, monkeypatch):
    """Файл можно приложить и к уже заведённой закупке — заказчик прислал уточнённое ТЗ."""

    _use_storage(monkeypatch, tmp_path)

    created = client.post(
        "/tenders/manual",
        headers=_auth_headers(admin_token),
        data={"title": "Заявка, файлы позже"},
    )
    tender_id = created.json()["tender"]["id"]

    response = client.post(
        f"/tenders/{tender_id}/documents/upload",
        headers=_auth_headers(admin_token),
        files=[
            ("files", ("ТЗ.docx", _docx_bytes(["Номинальное напряжение 230 В"]), "application/octet-stream")),
            ("files", ("Заметка.txt", b"\xd0\xa2\xd0\xb5\xd0\xba\xd1\x81\xd1\x82", "text/plain")),
        ],
    )
    assert response.status_code == 201, response.text
    documents = response.json()
    assert [d["file_name"] for d in documents] == ["ТЗ.docx", "Заметка.txt"]
    assert documents[0]["document_class"] == "tz_description"
    assert all(d["has_text"] for d in documents)

    listed = client.get(f"/tenders/{tender_id}/documents", headers=_auth_headers(admin_token))
    assert len(listed.json()) == 2

    db = SessionLocal()
    try:
        tender = db.get(Tender, uuid.UUID(tender_id))
        assert tender is not None and tender.source.key == MANUAL_SOURCE_KEY
    finally:
        db.close()


def test_upload_without_files_is_rejected(client, admin_token):
    created = client.post(
        "/tenders/manual", headers=_auth_headers(admin_token), data={"title": "Пустая"}
    )
    tender_id = created.json()["tender"]["id"]
    response = client.post(
        f"/tenders/{tender_id}/documents/upload",
        headers=_auth_headers(admin_token),
        files=[("files", ("", b"", "application/octet-stream"))],
    )
    assert response.status_code == 422
