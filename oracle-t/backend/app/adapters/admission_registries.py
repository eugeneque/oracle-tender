"""Выгрузки реестров допуска: ГИСП (реестр промышленной продукции, ПП РФ № 719) и перечень
оборудования, допущенного к применению на объектах ПАО «Россети» (замечание 13.10.2026, п.7:
«нет информации о реестрах допуска»; доступ к обоим настроен по замечанию 28.10.2026).

**ГИСП.** Таблица на странице реестра — грид DevExtreme, который берёт строки JSON-запросом
`POST /pp719v2/pub/prod/b/` с параметрами грида `{"opt": {skip, take, filter, …}}`. Тот же
запрос с фильтром по колонке «ОКПД2» (`_product_okpd2 startswith …`) отдаёт только нашу
продукцию: по 26.51.63 — около 700 записей против 490 тысяч строк полной выгрузки. Коды
ОКПД2 берутся из общих профилей отбора («наш ОКПД»), см. `registry_polling.our_okpd2_codes`.
Днём ГИСП отвечает минутами (замечание 28.10.2026), поэтому таймаут большой, есть повторы,
а плановый опрос стоит на ночь. Если запрос не прошёл, администратор загружает XLSX-выгрузку
(«Скачать только действующие») — её разбирает `parse_xlsx` с тем же фильтром по ОКПД2.

**Россети.** Перечень — PDF «Раздел II. Вторичное оборудование…» на странице «Проверка
качества оборудования»; адрес файла меняется с каждой редакцией («…на 22.10.2026.pdf»),
поэтому ссылка ищется на странице по подписи. Сайт закрыт проверкой браузера «Проверка
безопасности» (вычислительная задача на JavaScript) — её адаптер не проходит и не пытается:
если вместо страницы или PDF пришла проверка, это честная ошибка с подсказкой загрузить PDF
из браузера. Разбор PDF — по колонкам таблицы (`parse_pdf`).

Адаптер ничего не знает о каталоге: он превращает реестр в список строк `RegistryRow`
(номер, наименование продукции, организация, даты), а сопоставление с моделями — в
`registry_polling`.
"""

from __future__ import annotations

import io
import re
import time
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from typing import Iterable
from urllib.parse import urljoin

import httpx
from bs4 import BeautifulSoup
from loguru import logger

from app.adapters.http_utils import DEFAULT_USER_AGENT, resolve_verify

# Страница реестра ГИСП и JSON-запрос её грида.
GISP_PP719_PAGE = "https://gisp.gov.ru/pp719v2/pub/prod/"
GISP_PP719_API = "https://gisp.gov.ru/pp719v2/pub/prod/b/"
# Страница со ссылками на перечни допущенного оборудования (адрес из гайда 28.10.2026;
# прежний `/suppliers-and-contractors/attestation/` больше не существует).
ROSSETI_EQUIPMENT_PAGE = "https://www.rosseti.ru/suppliers/technical-policy/equipment-quality-control/"

REQUEST_TIMEOUT_SECONDS = 120.0
# Днём ГИСП отвечает на запрос грида по нескольку минут.
GISP_TIMEOUT_SECONDS = 300.0
GISP_ATTEMPTS = 3
GISP_RETRY_PAUSE_SECONDS = 30.0
# Грид отдаёт до 1000 строк за запрос; 500 — с запасом, чтобы медленный ответ не обрывался.
GISP_PAGE_SIZE = 500
# Полная выгрузка ГИСП — десятки мегабайт; всё заметно крупнее — не та ссылка.
MAX_FILE_BYTES = 300 * 1024 * 1024


class RegistryUnavailable(Exception):
    """Выгрузку получить не удалось. Текст — для человека: что произошло и что делать."""


@dataclass
class RegistryRow:
    record_number: str | None
    product_name: str
    organization: str = ""
    okpd2: str = ""
    issued_at: date | None = None
    valid_to: date | None = None
    # Текст, в котором ищется обозначение модели: наименование и организация, у PDF —
    # вся строка перечня.
    raw: str = ""
    # Ссылка на карточку записи (у ГИСП — карточка продукции), если реестр её даёт.
    url: str | None = None
    # Что ещё полезно тендерному отделу: баллы ПП 719, ТУ, баллы по СТО Россетей.
    details: str = ""
    # Баллы записи ГИСП по ПП 719 — по ним выбирается запись, годная для закупок.
    score: float | None = None


@dataclass
class RegistryFile:
    content: bytes
    file_name: str
    url: str | None = None
    # Строки уже разобраны (ответ JSON-запроса ГИСП) — разбирать `content` не нужно.
    rows: list[RegistryRow] | None = None
    total: int = 0


# --------------------------------------------------------------------------- скачивание


def _client(url: str, timeout: float = REQUEST_TIMEOUT_SECONDS) -> httpx.Client:
    return httpx.Client(
        headers={"User-Agent": DEFAULT_USER_AGENT, "Accept-Language": "ru-RU,ru;q=0.9"},
        timeout=timeout,
        follow_redirects=True,
        verify=resolve_verify(url),
    )


def is_browser_check(response: httpx.Response) -> bool:
    """Страница «Проверка безопасности» вместо содержимого: антибот-защита Россетей отдаёт
    её с кодом 200, и без этой проверки она разбиралась бы как пустая страница."""

    if "html" not in response.headers.get("content-type", ""):
        return False
    head = response.text[:3000]
    return "Проверка безопасности" in head or ('id="challenge"' in response.text and "nonce" in response.text)


def _get(client: httpx.Client, url: str) -> httpx.Response:
    try:
        response = client.get(url)
    except httpx.HTTPError as exc:
        raise RegistryUnavailable(f"{url} не ответил: {exc}") from exc
    if response.status_code == 403:
        raise RegistryUnavailable(
            f"{url} отказал роботу (403). Скачайте выгрузку в браузере и загрузите файл вручную."
        )
    if response.status_code >= 400:
        raise RegistryUnavailable(f"{url} ответил {response.status_code}")
    if len(response.content) > MAX_FILE_BYTES:
        raise RegistryUnavailable(f"{url}: файл больше {MAX_FILE_BYTES // (1024 * 1024)} МБ")
    return response


def find_download_link(html: str, base_url: str, *, extensions: tuple[str, ...], prefer: tuple[str, ...]) -> str | None:
    """Ссылка на выгрузку на странице: по расширению в адресе либо по подписи «скачать».
    Из нескольких — та, в подписи или адресе которой есть слово из `prefer` («вторичн»)."""

    soup = BeautifulSoup(html, "lxml")
    candidates: list[tuple[int, str]] = []
    for anchor in soup.find_all("a", href=True):
        href = anchor["href"].strip()
        text = " ".join(anchor.get_text(" ").split()).lower()
        lowered = href.lower()
        is_file = any(ext in lowered for ext in extensions)
        is_download = "скачать" in text or "выгруз" in text or anchor.has_attr("download")
        if not (is_file or is_download):
            continue
        score = (2 if is_file else 0) + (1 if is_download else 0)
        if any(word in text or word in lowered for word in prefer):
            score += 3
        candidates.append((score, urljoin(base_url, href)))
    if not candidates:
        return None
    candidates.sort(key=lambda item: -item[0])
    return candidates[0][1]


def okpd2_prefixes(codes: Iterable[str]) -> list[str]:
    """Коды без вложенных: из «26.51.63» и «26.51.63.130» достаточно первого."""

    cleaned = sorted({code.strip().rstrip(".") for code in codes if code and code.strip()}, key=lambda code: (len(code), code))
    result: list[str] = []
    for code in cleaned:
        if not any(code == kept or code.startswith(kept + ".") for kept in result):
            result.append(code)
    return result


def _gisp_filter(prefixes: list[str]) -> list:
    conditions: list = []
    for prefix in prefixes:
        if conditions:
            conditions.append("or")
        conditions.append(["_product_okpd2", "startswith", prefix])
    return conditions[0] if len(conditions) == 1 else conditions


def _gisp_post(client: httpx.Client, body: dict) -> dict:
    last_error = ""
    for attempt in range(1, GISP_ATTEMPTS + 1):
        try:
            response = client.post(GISP_PP719_API, json=body)
        except httpx.HTTPError as exc:
            last_error = f"не ответил: {exc.__class__.__name__}"
        else:
            if response.status_code == 403:
                raise RegistryUnavailable(
                    "ГИСП отказал серверу (403) — доступ к реестру открыт только из российских сетей. "
                    "Скачайте в браузере «Скачать только действующие» и загрузите XLSX."
                )
            if response.status_code < 400:
                try:
                    data = response.json()
                except ValueError:
                    raise RegistryUnavailable(
                        "ГИСП вместо данных реестра вернул страницу (проверка браузера или перенаправление). "
                        "Скачайте выгрузку XLSX в браузере и загрузите файл."
                    ) from None
                if not data.get("ok"):
                    raise RegistryUnavailable(f"ГИСП отклонил запрос реестра: {str(data)[:300]}")
                return data
            last_error = f"ответил {response.status_code}"
        logger.info(f"ГИСП: попытка {attempt} из {GISP_ATTEMPTS} — {last_error}")
        if attempt < GISP_ATTEMPTS:
            time.sleep(GISP_RETRY_PAUSE_SECONDS)
    raise RegistryUnavailable(
        f"ГИСП {last_error} после {GISP_ATTEMPTS} попыток. Днём реестр работает очень медленно — "
        "повторите вечером или загрузите XLSX-выгрузку из браузера."
    )


def _iso_date(value) -> date | None:
    if not value:
        return None
    try:
        return date.fromisoformat(str(value)[:10])
    except ValueError:
        return _to_date(value)


def _effective_valid_to(valid_till: date | None, ended_at: date | None) -> date | None:
    """Запись, действие которой прекращено досрочно, действует до даты прекращения."""

    if ended_at and (valid_till is None or ended_at < valid_till):
        return ended_at
    return valid_till


def _score(value) -> float | None:
    try:
        return float(str(value).replace(",", ".")) if value not in (None, "") else None
    except ValueError:
        return None


def _gisp_details(score, percentage, spec) -> str:
    parts = []
    if score not in (None, ""):
        parts.append(f"баллов по ПП 719: {score}")
    if percentage not in (None, ""):
        parts.append(f"процентный показатель: {percentage}")
    if spec:
        parts.append(f"изготовлена по {spec}")
    return "; ".join(parts)


def gisp_item_row(item: dict) -> RegistryRow:
    """Строка ответа грида ГИСП → `RegistryRow`. Поля — как `dataField` колонок грида."""

    product = (item.get("_product_name") or "").strip()
    organization = (item.get("_org_name") or "").strip()
    return RegistryRow(
        record_number=(
            item.get("_product_reg_number_2023")
            or item.get("_product_reg_number_2022")
            or item.get("_res_number")
            or None
        ),
        product_name=product,
        organization=organization,
        okpd2=(item.get("_product_okpd2") or "").strip(),
        issued_at=_iso_date(item.get("_res_date")),
        valid_to=_effective_valid_to(_iso_date(item.get("_res_valid_till")), _iso_date(item.get("_res_end_date"))),
        raw=" ".join(part for part in (product, organization) if part),
        url=item.get("_product_gisp_url") or item.get("_gisp_url") or None,
        details=_gisp_details(
            item.get("_product_score_value"), item.get("_product_percentage"), (item.get("_product_spec") or "").strip()
        ),
        score=_score(item.get("_product_score_value")),
    )


def fetch_gisp(okpd2_codes: Iterable[str]) -> RegistryFile:
    """Записи реестра ПП 719 по нашим кодам ОКПД2 — запросом грида, страницами."""

    prefixes = okpd2_prefixes(okpd2_codes)
    if not prefixes:
        raise RegistryUnavailable("Не заданы коды ОКПД2 для отбора в ГИСП — заполните их в общих профилях отбора")
    rows: list[RegistryRow] = []
    total = None
    with _client(GISP_PP719_API, timeout=GISP_TIMEOUT_SECONDS) as client:
        client.headers.update({"Accept": "application/json", "Referer": GISP_PP719_PAGE})
        skip = 0
        while total is None or skip < total:
            data = _gisp_post(
                client,
                {
                    "opt": {
                        "skip": skip,
                        "take": GISP_PAGE_SIZE,
                        "requireTotalCount": True,
                        "filter": _gisp_filter(prefixes),
                        "sort": [{"selector": "_product_reg_number_2023", "desc": False}],
                    }
                },
            )
            items = data.get("items") or []
            total = int(data.get("total_count") or 0) if total is None else total
            rows.extend(gisp_item_row(item) for item in items)
            if not items:
                break
            skip += len(items)
    return RegistryFile(
        content=b"",
        file_name=f"ГИСП, ОКПД2 {', '.join(prefixes)}",
        url=GISP_PP719_PAGE,
        rows=rows,
        total=total or len(rows),
    )


def download_rosseti(known_url: str | None = None) -> RegistryFile:
    """PDF перечня Россетей (раздел II, вторичное оборудование). Ищется на странице
    «Проверка качества оборудования» по подписи; если не нашёлся — последний известный адрес."""

    with _client(ROSSETI_EQUIPMENT_PAGE) as client:
        link: str | None = None
        challenged = False
        try:
            page = _get(client, ROSSETI_EQUIPMENT_PAGE)
            if is_browser_check(page):
                challenged = True
            else:
                link = find_download_link(
                    page.text,
                    ROSSETI_EQUIPMENT_PAGE,
                    extensions=(".pdf",),
                    prefer=("вторичн", "раздел ii", "раздел%20ii"),
                )
        except RegistryUnavailable as exc:
            logger.info(f"Перечень Россетей: страница недоступна ({exc}), пробую известный адрес")
        link = link or known_url
        if not link:
            raise RegistryUnavailable("Адрес перечня ПАО «Россети» неизвестен — загрузите PDF вручную")
        response = _get(client, link)
        if not response.content.startswith(b"%PDF"):
            reason = (
                "сайт Россетей закрыт проверкой браузера «Проверка безопасности»"
                if challenged or is_browser_check(response)
                else f"по адресу {link} пришёл не PDF"
            )
            raise RegistryUnavailable(
                f"Перечень не скачан: {reason}. Откройте {ROSSETI_EQUIPMENT_PAGE}, скачайте "
                "«Раздел II. Вторичное оборудование…» и загрузите PDF."
            )
        return RegistryFile(response.content, link.rsplit("/", 1)[-1] or "rosseti.pdf", link)


# --------------------------------------------------------------------------- разбор XLSX

_DATE_RE = re.compile(r"(\d{2})\.(\d{2})\.(\d{4})")


def _to_date(value) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and 20000 < value < 80000:
        # Дата Excel числом дней от 30.12.1899.
        return date(1899, 12, 30) + timedelta(days=int(value))
    match = _DATE_RE.search(str(value))
    if match:
        day, month, year = (int(part) for part in match.groups())
        try:
            return date(year, month, day)
        except ValueError:
            return None
    return None


def _column_roles(header: list[str]) -> dict[str, int]:
    """Колонки выгрузки по словам заголовка. Шапка ГИСП на 26.10.2026: «Предприятие», «ИНН»,
    «Первичный регистрационный номер реестровой записи» (почти всегда пустой), «Реестровый
    номер», «Дата внесения в реестр», «Срок действия», «Фактическая дата прекращения…»,
    «Наименование продукции», «ОКПД2», «Изготовлена по», «Баллы», «Процентный показатель»."""

    roles: dict[str, int] = {}
    for index, raw in enumerate(header):
        text = (raw or "").lower().replace("ё", "е")
        if not text:
            continue
        if "номер" in text and ("реестр" in text or "запис" in text):
            # Первичный номер записи (до 2023 года) — запасной: у новых записей он пуст.
            roles.setdefault("number_primary" if "первичн" in text else "number", index)
        elif "окпд" in text and "okpd2" not in roles:
            roles["okpd2"] = index
        elif text.strip() == "инн" or text.startswith("инн "):
            roles.setdefault("inn", index)
        elif "наименование" in text and any(w in text for w in ("продукц", "товар", "изделия")):
            roles.setdefault("product", index)
        elif "прекращ" in text:
            roles.setdefault("ended_at", index)
        elif any(w in text for w in ("организац", "предприят", "производител", "изготовител", "юридическ")):
            roles.setdefault("organization", index)
        elif "дата" in text and any(w in text for w in ("внесен", "включ", "регистрац", "выдач")):
            roles.setdefault("issued_at", index)
        elif ("действ" in text and ("до" in text or "срок" in text or "окончан" in text)) or "окончани" in text:
            roles.setdefault("valid_to", index)
        elif text.strip() == "баллы":
            roles.setdefault("score", index)
        elif "процентн" in text:
            roles.setdefault("percentage", index)
        elif text.startswith("изготовлен"):
            roles.setdefault("spec", index)
    if "number" not in roles and "number_primary" in roles:
        roles["number"] = roles.pop("number_primary")
    return roles


def _okpd2_matches(okpd2: str, prefixes: list[str]) -> bool:
    return any(okpd2 == prefix or okpd2.startswith(prefix + ".") for prefix in prefixes)


def parse_xlsx(
    content: bytes, *, relevant: re.Pattern[str], okpd2_codes: Iterable[str] = ()
) -> tuple[list[RegistryRow], int]:
    """Строки выгрузки ГИСП. Возвращаются только наши: код ОКПД2 из `okpd2_codes` (коды
    общих профилей отбора) либо совпадение с `relevant` (счётчики электроэнергии по
    наименованию), и общее число строк.

    Полная выгрузка — около 500 тысяч строк, поэтому в текст переводятся только нужные
    колонки: перевод всех 28 колонок каждой строки занимал около минуты."""

    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - битый файл → понятная ошибка, а не трассировка
        raise RegistryUnavailable(f"Файл не читается как таблица XLSX: {exc}") from exc

    prefixes = okpd2_prefixes(okpd2_codes)
    rows: list[RegistryRow] = []
    total = 0
    for sheet in workbook.worksheets:
        roles: dict[str, int] | None = None
        for values in sheet.iter_rows(values_only=True):
            if roles is None:
                candidate = _column_roles(["" if v is None else str(v).strip() for v in values])
                if "number" in candidate and "product" in candidate:
                    roles = candidate
                continue
            if not any(v not in (None, "") for v in values):
                continue
            total += 1

            def raw_cell(role: str):
                index = roles.get(role)
                return values[index] if index is not None and index < len(values) else None

            def cell(role: str) -> str:
                value = raw_cell(role)
                return "" if value is None else str(value).strip()

            okpd2 = cell("okpd2")
            product = cell("product")
            if not (_okpd2_matches(okpd2, prefixes) or relevant.search(okpd2) or relevant.search(product)):
                continue
            organization = cell("organization")
            rows.append(
                RegistryRow(
                    record_number=cell("number") or cell("number_primary") or None,
                    product_name=product,
                    organization=organization,
                    okpd2=okpd2,
                    issued_at=_to_date(raw_cell("issued_at")),
                    valid_to=_effective_valid_to(_to_date(raw_cell("valid_to")), _to_date(raw_cell("ended_at"))),
                    raw=" ".join(part for part in (product, organization) if part),
                    details=_gisp_details(raw_cell("score"), raw_cell("percentage"), cell("spec")),
                    score=_score(raw_cell("score")),
                )
            )
    workbook.close()
    if total == 0:
        raise RegistryUnavailable(
            "В таблице не нашлось строки заголовка с номером реестровой записи и наименованием "
            "продукции — это не выгрузка реестра ПП 719?"
        )
    return rows, total


# --------------------------------------------------------------------------- разбор PDF
#
# Перечень Россетей — таблица из шести колонок: «№ п/п», «Производитель/Заявитель»,
# «Наименование оборудования», «Дата утверждения ЗАК», «Срок действия ЗАК / ИПП», «№ ЗАК,
# протокола продления». Строка таблицы занимает от двух строк текста до нескольких страниц,
# номер в ячейке бывает выровнен по середине (выше него — текст той же записи), а у подряд
# идущих моделей одного производителя ячейки производителя и ЗАК объединены и текст стоит
# только в одной из строк. Поэтому границы записей берутся не из текста, а из линий
# таблицы: тонких горизонтальных прямоугольников под колонкой номера. Проверено на редакции
# 22.10.2026 (110 страниц, 245 записей): все номера идут подряд, у 242 есть номер ЗАК.

_ZAK_NUMBER_RE = re.compile(r"\b[IVІ]{1,3}[А-ЯA-Z]{1,3}-\d{1,4}/\d{2}\b")
_ZAK_FALLBACK_RE = re.compile(r"(?:ЗАК|заключени\w*)[^№\n]{0,30}№\s*([\w./-]+)", re.IGNORECASE)
_STO_SCORE_RE = re.compile(r"(\d+(?:[.,]\d+)?)\s*балл", re.IGNORECASE)
_ROW_NUMBER_RE = re.compile(r"^\d{1,4}$")

# Сегмент текста: левый край, правый край, текст.
Segment = tuple[float, float, str]


@dataclass
class _PdfPage:
    width: float
    # Строки текста сверху вниз: (верх строки, сегменты слева направо).
    lines: list[tuple[float, list[Segment]]]
    # Верхние края горизонтальных линий таблицы, пересекающих колонку номера.
    separators: list[float]


@dataclass
class _PdfRecord:
    number: str
    segments: list[list[Segment]] = field(default_factory=list)


def _pdf_pages(content: bytes) -> list[_PdfPage]:
    import pdfplumber

    pages: list[_PdfPage] = []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for page in pdf.pages:
            words = page.extract_words(keep_blank_chars=True, x_tolerance=2, y_tolerance=2, use_text_flow=False)
            words.sort(key=lambda w: (round(w["top"]), w["x0"]))
            grouped: list[tuple[float, list[dict]]] = []
            for word in words:
                if grouped and abs(grouped[-1][0] - word["top"]) <= 2.5:
                    grouped[-1][1].append(word)
                else:
                    grouped.append((word["top"], [word]))
            lines: list[tuple[float, list[Segment]]] = []
            for top, line_words in grouped:
                segments: list[Segment] = []
                for word in sorted(line_words, key=lambda w: w["x0"]):
                    text = word["text"].strip()
                    if not text:
                        continue
                    # Слова одной ячейки идут вплотную; разрыв шире 6 pt — уже другая колонка.
                    if segments and word["x0"] - segments[-1][1] < 6:
                        x0, _, prev = segments[-1]
                        segments[-1] = (x0, word["x1"], f"{prev} {text}")
                    else:
                        segments.append((word["x0"], word["x1"], text))
                if segments:
                    lines.append((top, segments))
            separators = sorted(
                {
                    round(obj["top"], 1)
                    for obj in list(page.rects) + list(page.lines)
                    if abs(obj["bottom"] - obj["top"]) <= 2 and obj["x0"] < 30 < obj["x1"]
                }
            )
            pages.append(_PdfPage(width=float(page.width), lines=lines, separators=separators))
    return pages


def _row_number(segments: list[Segment], width: float) -> str | None:
    for x0, _, text in segments:
        if x0 < width * 0.06 and _ROW_NUMBER_RE.match(text):
            return text
    return None


def _pdf_records(pages: list[_PdfPage]) -> list[_PdfRecord]:
    """Записи перечня: строки текста, разложенные по полосам между линиями таблицы. Полоса
    без номера — продолжение записи с прошлой страницы. На странице без линий (другая
    вёрстка) запись начинается со строки, где слева стоит номер."""

    records: list[_PdfRecord] = []
    current: _PdfRecord | None = None
    for page in pages:
        if len(page.separators) >= 2:
            bands: list[list[list[Segment]]] = [[] for _ in range(len(page.separators) - 1)]
            for top, segments in page.lines:
                for index in range(len(page.separators) - 1):
                    if page.separators[index] < top < page.separators[index + 1]:
                        bands[index].append(segments)
                        break
            for band in bands:
                number = next((n for n in (_row_number(s, page.width) for s in band) if n), None)
                if number:
                    current = _PdfRecord(number=number)
                    records.append(current)
                if current is not None:
                    current.segments.extend(band)
            continue
        for _, segments in page.lines:
            number = _row_number(segments, page.width)
            if number and len(segments) > 1:
                current = _PdfRecord(number=number)
                records.append(current)
            elif current is not None and len(segments) == 1 and _ROW_NUMBER_RE.match(segments[0][2]):
                continue  # номер страницы
            if current is not None:
                current.segments.append(segments)
    return records


def _column_starts(records: list[_PdfRecord], width: float) -> tuple[float, float, float]:
    """Левые края колонок наименования, даты утверждения и срока действия — самые частые
    позиции у строк, где они стоят рядом с номером записи."""

    name_x: Counter[int] = Counter()
    issued_x: Counter[int] = Counter()
    valid_x: Counter[int] = Counter()
    for record in records:
        for segments in record.segments:
            text_segments = [s for s in segments if not _DATE_RE.fullmatch(s[2]) and not _ROW_NUMBER_RE.match(s[2])]
            if len(text_segments) >= 2 and text_segments[1][0] < width * 0.7:
                name_x[round(text_segments[1][0])] += 1
            dates = [s for s in segments if _DATE_RE.fullmatch(s[2])]
            if len(dates) >= 2:
                issued_x[round(dates[0][0])] += 1
                valid_x[round(dates[1][0])] += 1

    def mode(counter: Counter[int], default: float) -> float:
        return float(counter.most_common(1)[0][0]) if counter else default

    return mode(name_x, width * 0.22), mode(issued_x, width * 0.72), mode(valid_x, width * 0.79)


def _join_lines(lines: list[str]) -> str:
    """Строки ячейки в одну: перенос по дефису («МИРТЕК-» + «135-РУ») склеивается без пробела."""

    text = ""
    for line in lines:
        if text.endswith("-") and not text.endswith(" -"):
            text += line
        else:
            text = f"{text} {line}" if text else line
    return text


def _rows_from_records(records: list[_PdfRecord], width: float) -> list[RegistryRow]:
    name_x, issued_x, valid_x = _column_starts(records, width)
    date_boundary = (issued_x + valid_x) / 2
    rows: list[RegistryRow] = []
    previous: RegistryRow | None = None
    for record in records:
        manufacturer: list[str] = []
        name: list[str] = []
        right: list[str] = []
        issued: list[date] = []
        valid: list[date] = []
        zak_main: str | None = None
        for segments in record.segments:
            if len(segments) == 1 and _ROW_NUMBER_RE.match(segments[0][2]):
                continue  # номер записи или номер страницы
            for x0, _, text in segments:
                if x0 < width * 0.06 and _ROW_NUMBER_RE.match(text):
                    continue
                if x0 >= issued_x - 8:
                    right.append(text)
                    for match in _DATE_RE.finditer(text):
                        value = _to_date(match.group(0))
                        if value:
                            (issued if x0 < date_boundary else valid).append(value)
                    if zak_main is None and (found := _ZAK_NUMBER_RE.search(text)):
                        zak_main = found.group(0)
                elif x0 >= name_x - 10:
                    name.append(text)
                else:
                    manufacturer.append(text)
        organization = _join_lines(manufacturer)
        product = _join_lines(name)
        if not product:
            continue
        right_text = " ".join(right)
        if zak_main is None and (fallback := _ZAK_FALLBACK_RE.search(right_text)):
            zak_main = fallback.group(1)
        inherited = False
        if not organization and previous is not None:
            # Объединённая ячейка: производитель (и обычно ЗАК) — у соседней записи.
            organization = previous.organization
            inherited = True
        if not organization:
            continue  # заголовок раздела («АИИС КУЭ»), а не запись
        row = RegistryRow(
            record_number=zak_main,
            product_name=product,
            organization=organization,
            issued_at=min(issued) if issued else None,
            valid_to=max(valid) if valid else None,
            raw=f"{product} {organization}",
        )
        if inherited and row.record_number is None and row.valid_to is None:
            row.record_number = previous.record_number
            row.issued_at = previous.issued_at
            row.valid_to = previous.valid_to
        details = [f"№ п/п {record.number} в перечне"]
        if inherited:
            details.append("производитель и ЗАК — из объединённой ячейки предыдущей строки")
        score = _STO_SCORE_RE.search(product)
        if score:
            details.append(f"баллов по СТО 34.01-22-002: {score.group(1)}")
        row.details = "; ".join(details)
        rows.append(row)
        previous = row
    return rows


def parse_pdf(content: bytes) -> tuple[list[RegistryRow], int]:
    """Записи перечня Россетей и их число. Если таблицу разобрать не удалось (другая вёрстка
    PDF), строкой считается абзац текста, а номер ЗАК и срок ищутся в нём и двух следующих."""

    try:
        pages = _pdf_pages(content)
    except Exception as exc:  # noqa: BLE001 - см. parse_xlsx
        raise RegistryUnavailable(f"PDF не читается: {exc}") from exc
    if not any(page.lines for page in pages):
        raise RegistryUnavailable("В PDF нет текстового слоя (скан?) — сверка возможна только вручную")
    width = max((page.width for page in pages), default=842.0)
    rows = _rows_from_records(_pdf_records(pages), width)
    if rows:
        return rows, len(rows)
    logger.info("Перечень Россетей: таблица не распознана, разбираю PDF построчно")
    return _parse_pdf_lines([" ".join(s[2] for s in segments) for page in pages for _, segments in page.lines])


def _parse_pdf_lines(lines: list[str]) -> tuple[list[RegistryRow], int]:
    rows: list[RegistryRow] = []
    for index, line in enumerate(lines):
        window = " ".join(lines[index : index + 3])
        number = _ZAK_NUMBER_RE.search(window) or _ZAK_FALLBACK_RE.search(window)
        dates = [_to_date(match.group(0)) for match in _DATE_RE.finditer(window)]
        dates = [item for item in dates if item]
        rows.append(
            RegistryRow(
                record_number=(number.group(1) if number.groups() else number.group(0)) if number else None,
                product_name=line,
                valid_to=max(dates) if dates else None,
                raw=window,
            )
        )
    return rows, len(lines)
