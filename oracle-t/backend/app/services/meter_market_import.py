"""Импорт справочника рынка ПУ из Excel аналитика (раздел 79 ARCHITECTURE.md, раздел 6.2 ТЗ).

Стартовые данные — файл «3 кв 2026.xlsx», лист «Лист1»: строка = тип ГРСИ × **сырое
написание** модификации из поверок с атрибутами отчёта (производитель, 1ф/3ф, исполнение,
интерфейс, МПИ, локализация) и количеством за квартал. То есть это уже разметка вариантов:
каждая строка становится вариантом, а строки одного типа с одним ключом нормализации и
одинаковыми атрибутами — одной модификацией.

Данные грязные (раздел 6.2 ТЗ): «шкафное»/«Шкафное», « да» с пробелом, срок свидетельства то
датой, то строкой, табуляции в названиях, «н/д» и «-» в числовых полях. Лист растянут до
1 038 697 строк, из них значимых 1 194 — чтение останавливается на длинной полосе пустых строк.

Повторный импорт безопасен: ручная разметка (`mapped_by = manual`) и ручная локализация
(`localization_source = manual`) не перезаписываются.

Запуск: `python -m app.services.meter_market_import "<путь к xlsx>"`.
"""

from __future__ import annotations

import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import IO, Any

import openpyxl
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.meter_market import (
    LocalizationSource,
    MappedBy,
    MeterMarketGroup,
    MeterMarketGroupMember,
    MeterMarketManufacturer,
    MeterMarketModification,
    MeterMarketType,
    MeterMarketVariant,
    TypeReviewStatus,
)
from app.services.meter_market_normalize import (
    EMPTY_MODIFICATION,
    clean_text,
    manufacturer_display_name,
    manufacturer_key,
    modification_key,
    normalize_design,
    normalize_interface,
    normalize_phase,
)

SHEET = "Лист1"
_EMPTY_RUN_STOP = 2000

# Группы компаний по отчёту аналитика за 3 кв. 2026 (раздел 6.4 ТЗ): ключ производителя
# (`manufacturer_key`) и годы участия включительно.
DEFAULT_GROUPS: dict[str, list[tuple[str, int | None, int | None]]] = {
    "Завод НАРТИС": [("заводнартис", None, None), ("атомцифросбыт", 2026, None)],
    "Энергомера": [("энергомера", None, None), ("атомцифросбыт", 2022, 2025)],
    "МИРТЕК": [
        ("миртек", None, None),
        ("умныйучет", None, None),
        ("каскад", None, None),
        ("оэс", None, None),
        ("электролабораторияэлпэк", None, None),
    ],
    "НПО «МИР»": [("нпомир", None, None), ("ргмэк", None, None)],
}


@dataclass
class ImportRow:
    manufacturer: str
    mit_number: str
    cert_valid_to: date | None
    phase: str | None
    design: str
    count: int | None
    interface: str | None
    mpi: int | None
    notation: str | None
    raw: str
    comment: str | None
    in_rpp: bool
    score: float | None


@dataclass
class ImportStats:
    rows: int = 0
    quantity: int = 0
    types: int = 0
    manufacturers: int = 0
    modifications: int = 0
    variants: int = 0
    skipped_manual: int = 0
    warnings: list[str] = field(default_factory=list)


def _as_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = clean_text(value)
    if not text:
        return None
    try:
        return datetime.strptime(text[:10], "%d.%m.%Y").date()
    except ValueError:
        return None


def _as_number(value: Any) -> float | None:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    text = (clean_text(value) or "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def _as_raw(value: Any) -> str:
    """Сырое написание модификации — как в поверке, без чистки: по нему вариант ищется в
    выгрузке ФГИС. Числа Excel (обозначение «107») — без «.0»."""

    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return "" if value is None else str(value)


def read_rows(source: str | Path | IO[bytes]) -> list[ImportRow]:
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = workbook[SHEET] if SHEET in workbook.sheetnames else workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        header = [clean_text(h) or "" for h in next(rows)]
        col = {name: i for i, name in enumerate(header)}
        required = ("Производитель", "Регистрационный номер типа СИ", "Модификация СИ")
        missing = [name for name in required if name not in col]
        if missing:
            raise ValueError(f"В листе нет колонок: {', '.join(missing)}")

        def get(row: tuple, name: str) -> Any:
            index = col.get(name)
            return row[index] if index is not None and index < len(row) else None

        result: list[ImportRow] = []
        empty_run = 0
        for row in rows:
            if not any(v not in (None, "") for v in row):
                empty_run += 1
                if empty_run >= _EMPTY_RUN_STOP:
                    break
                continue
            empty_run = 0
            number = clean_text(get(row, "Регистрационный номер типа СИ"))
            manufacturer = clean_text(get(row, "Производитель"))
            if not number or not manufacturer:
                continue
            count = _as_number(get(row, "Количество"))
            mpi = _as_number(get(row, "Межповерочный интервал"))
            result.append(
                ImportRow(
                    manufacturer=manufacturer,
                    mit_number=number,
                    cert_valid_to=_as_date(get(row, "Срок свидетельства")),
                    phase=normalize_phase(get(row, "Тип ПУ")),
                    design=normalize_design(get(row, "Исполнение")),
                    count=int(count) if count is not None else None,
                    interface=normalize_interface(get(row, "Интерфейс")),
                    mpi=int(mpi) if mpi is not None else None,
                    notation=clean_text(_as_raw(get(row, "Наименование типа СИ"))),
                    raw=_as_raw(get(row, "Модификация СИ")),
                    comment=clean_text(get(row, "Дополнительно")),
                    in_rpp=(clean_text(get(row, "В Реестре российской промышленной продукции")) or "").casefold()
                    == "да",
                    score=_as_number(get(row, "Кол-во баллов")),
                )
            )
        return result
    finally:
        workbook.close()


def _manufacturer(db: Session, cache: dict[str, MeterMarketManufacturer], raw: str, stats: ImportStats) -> MeterMarketManufacturer:
    key = manufacturer_key(raw)
    found = cache.get(key)
    if found is not None:
        return found
    name = manufacturer_display_name(raw)
    found = db.scalar(select(MeterMarketManufacturer).where(MeterMarketManufacturer.name == name))
    if found is None:
        found = MeterMarketManufacturer(name=name, synonyms=[key])
        db.add(found)
        db.flush()
        stats.manufacturers += 1
    elif key not in found.synonyms:
        found.synonyms = [*found.synonyms, key]
    cache[key] = found
    return found


def _ensure_groups(db: Session, cache: dict[str, MeterMarketManufacturer]) -> None:
    for group_name, members in DEFAULT_GROUPS.items():
        group = db.scalar(select(MeterMarketGroup).where(MeterMarketGroup.name == group_name))
        if group is not None:
            continue  # группу уже правили — импорт состав не трогает
        group = MeterMarketGroup(name=group_name)
        db.add(group)
        db.flush()
        for key, year_from, year_to in members:
            manufacturer = cache.get(key)
            if manufacturer is not None:
                db.add(
                    MeterMarketGroupMember(
                        group_id=group.id, manufacturer_id=manufacturer.id, year_from=year_from, year_to=year_to
                    )
                )


def import_rows(db: Session, rows: list[ImportRow]) -> ImportStats:
    stats = ImportStats(rows=len(rows), quantity=sum(r.count or 0 for r in rows))
    manufacturers: dict[str, MeterMarketManufacturer] = {}
    for manufacturer in db.scalars(select(MeterMarketManufacturer)):
        for key in manufacturer.synonyms or []:
            manufacturers[key] = manufacturer

    by_type: dict[str, list[ImportRow]] = defaultdict(list)
    for row in rows:
        by_type[row.mit_number].append(row)

    for number, type_rows in by_type.items():
        # Владелец типа — производитель большинства строк (правило 1 ТЗ); у 61891-15 строки
        # записаны и на МИРТЕК, и на «Умный учёт» — остальные уходят переопределением модификации.
        owner_raw = Counter(r.manufacturer for r in type_rows).most_common(1)[0][0]
        owner = _manufacturer(db, manufacturers, owner_raw, stats)
        si_type = db.scalar(select(MeterMarketType).where(MeterMarketType.mit_number == number))
        if si_type is None:
            si_type = MeterMarketType(mit_number=number)
            db.add(si_type)
            stats.types += 1
        si_type.manufacturer_id = owner.id
        si_type.in_report = True
        si_type.review_status = TypeReviewStatus.INCLUDED.value
        si_type.cert_valid_to = next((r.cert_valid_to for r in type_rows if r.cert_valid_to), si_type.cert_valid_to)
        si_type.mpi_years = Counter(r.mpi for r in type_rows if r.mpi).most_common(1)[0][0] if any(
            r.mpi for r in type_rows
        ) else si_type.mpi_years
        si_type.notation = si_type.notation or next((r.notation for r in type_rows if r.notation), None)
        db.flush()

        existing_mods = {m.name: m for m in db.scalars(select(MeterMarketModification).where(MeterMarketModification.type_id == si_type.id))}
        existing_variants = {v.raw: v for v in db.scalars(select(MeterMarketVariant).where(MeterMarketVariant.type_id == si_type.id))}

        # Дубли «тип × написание» в файле (строка с количеством и та же без него) — одна,
        # предпочтительно с количеством.
        unique: dict[str, ImportRow] = {}
        for row in type_rows:
            if row.raw not in unique or (unique[row.raw].count is None and row.count is not None):
                unique[row.raw] = row

        groups: dict[tuple, list[ImportRow]] = defaultdict(list)
        for row in unique.values():
            override = _manufacturer(db, manufacturers, row.manufacturer, stats)
            attrs = (
                row.phase,
                row.design,
                row.interface,
                row.mpi,
                row.in_rpp,
                row.score,
                None if override.id == owner.id else override.id,
            )
            groups[(modification_key(row.raw), attrs)].append(row)

        for (_, attrs), members in groups.items():
            phase, design, interface, mpi, in_rpp, score, override_id = attrs
            name = clean_text(max(members, key=lambda r: r.count or 0).raw) or EMPTY_MODIFICATION
            modification = existing_mods.get(name)
            if modification is not None and (modification.phase, modification.design, modification.interface) != (
                phase,
                design,
                interface,
            ) and modification.localization_source != LocalizationSource.MANUAL.value:
                # Одно имя, разные атрибуты — разные исполнения в справочнике аналитика.
                suffix = 2
                while f"{name} ({suffix})" in existing_mods:
                    suffix += 1
                name = f"{name} ({suffix})"
                modification = None
                stats.warnings.append(f"{number}: «{name}» — у одного написания разные атрибуты")
            if modification is None:
                modification = MeterMarketModification(type_id=si_type.id, name=name)
                db.add(modification)
                existing_mods[name] = modification
                stats.modifications += 1
            modification.phase = phase
            modification.design = design
            modification.interface = interface
            modification.mpi_years = mpi
            modification.manufacturer_override_id = override_id
            modification.comment = next((r.comment for r in members if r.comment), modification.comment)
            if modification.localization_source != LocalizationSource.MANUAL.value:
                modification.in_rpp = in_rpp
                modification.score = score if in_rpp else None
                modification.localization_source = LocalizationSource.IMPORT.value
            db.flush()

            for row in members:
                variant = existing_variants.get(row.raw)
                if variant is None:
                    variant = MeterMarketVariant(type_id=si_type.id, raw=row.raw, norm_key=modification_key(row.raw))
                    db.add(variant)
                    existing_variants[row.raw] = variant
                    stats.variants += 1
                elif variant.mapped_by == MappedBy.MANUAL.value:
                    stats.skipped_manual += 1
                    continue
                variant.modification_id = modification.id
                variant.mapped_by = MappedBy.IMPORT.value
        db.flush()

    _ensure_groups(db, manufacturers)
    db.flush()
    return stats


def import_file(db: Session, source: str | Path | IO[bytes]) -> ImportStats:
    return import_rows(db, read_rows(source))


if __name__ == "__main__":
    from app.db.session import SessionLocal

    session = SessionLocal()
    try:
        result = import_file(session, sys.argv[1])
        session.commit()
        print(result)
    finally:
        session.close()
