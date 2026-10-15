"""Учёт расхода токенов ИИ и «Детализация ИИ» для администратора (14.10.2026).

Руководство спросило, сколько Сова тратит на модель и какое окно контекста нужно ускорителю
с локальной моделью. Ответить «на глаз» нельзя: расход зависит от числа закупок, доли с
документацией, числа требований и производителей. Зато каждый провайдер возвращает точный
счётчик токенов в ответе — клиенты моделей (`routerai_client`, `yandex_ai_client`,
`gigachat_client`) передают его сюда после каждой попытки, а «Интеграции → Детализация ИИ»
сводит записи по дням, моделям, шагам разбора и закупкам.

Запись идёт **отдельной сессией** на том же подключении, что и у вызывающего кода
(`db.get_bind()`): расчёт матрицы может откатить свою транзакцию после сбоя на
производителе, а потраченные токены от этого не возвращаются. В тестах подключение — общее
с откатываемой транзакцией теста, и записи пропадают вместе с ней. Сбой записи учёта
никогда не роняет обращение к модели — только предупреждение в журнал.
"""

from __future__ import annotations

import io
import uuid
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal
from typing import Any

from loguru import logger
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import Date, case, cast, func, select
from sqlalchemy.orm import Session

from app.core.timezones import MOSCOW_TZ
from app.models.ai_usage import AiUsageRecord
from app.models.job import BackgroundJob
from app.models.log import LogLevel
from app.models.tender import Tender
from app.models.user import User
from app.services.ai_context import current_job_id, current_user_id
from app.services.ai_log_service import STATUS_ACTIONS, log_ai_event
from app.services.ai_provider_service import CHARS_PER_TOKEN, PROVIDER_LABELS

# Шаг разбора по схеме ответа. Неизвестная схема показывается своим именем — новый вызов
# модели не пропадёт из детализации, просто будет подписан по-английски, пока его не внесут.
OPERATION_LABELS: dict[str, str] = {
    "ClassificationResult": "Классификация документов",
    "RequirementsResult": "Извлечение требований",
    "ConsolidationResult": "Сведение требований",
    "ComplianceResult": "Матрица соответствия",
    "ExtractionResult": "Характеристики продукции",
    "SectionsBatchOne": "Разделы «Дополнительно»",
    "SectionsBatchTwo": "Разделы «Дополнительно»",
    "TaskAnswer": "AI-оценка по профилю",
    "CompetenciesAnswer": "AI-оценка по профилю",
    "DecisionAnswer": "AI-оценка по профилю",
    "ResumeAnswer": "AI-оценка по профилю",
    "FitAnswer": "Заключение ИИ",
    "PlanAnswer": "Заключение ИИ",
    "RelevanceAnswer": "ИИ-отбор закупок",
    "AiSummary": "Сводка аналитики",
}

JOB_KIND_LABELS: dict[str, str] = {
    "tender_analysis": "Анализ документов",
    "tender_evaluation": "Расчёт соответствия",
    "ai_profile_score": "AI-оценка",
    "tender_full_review": "Полный разбор закупки",
    "ai_feedback": "Пересмотр по замечанию",
}
# Обращения вне фоновой задачи: ИИ-отбор по расписанию, каталог, кнопки без очереди.
NO_JOB_LABEL = "Вне очереди разборов (расписание, каталог)"

STATUS_LABELS: dict[str, str] = {
    "success": "успешно",
    "invalid": "ответ не по схеме",
    "error": "сбой",
    "quota": "лимит исчерпан",
    "filter": "отказ фильтра",
}

MAX_ERROR_CHARS = 300
_SUM_FIELDS = (
    "calls",
    "prompt_tokens",
    "completion_tokens",
    "total_tokens",
    "cached_tokens",
    "failed",
    "duration_ms",
    "success_duration_ms",
    "success_completion_tokens",
)


@dataclass
class CallUsage:
    """Счётчик одного ответа провайдера, приведённый к общему виду."""

    prompt_tokens: int | None = None
    completion_tokens: int | None = None
    cached_tokens: int = 0
    cost: Decimal | None = None


def usage_from_openai(payload: dict[str, Any] | None) -> CallUsage:
    """`usage` OpenAI-совместимого ответа (RouterAI, GigaChat, Yandex в OpenAI-режиме).
    Кэш — `prompt_tokens_details.cached_tokens` (OpenAI, RouterAI) или
    `precached_prompt_tokens` (GigaChat); стоимость — если шлюз её вернул."""

    usage = (payload or {}).get("usage") or {}
    if not isinstance(usage, dict):
        return CallUsage()
    details = usage.get("prompt_tokens_details") or {}
    cached = (details.get("cached_tokens") if isinstance(details, dict) else None) or usage.get(
        "precached_prompt_tokens"
    )
    cost = usage.get("cost", usage.get("total_cost"))
    try:
        cost_value = Decimal(str(cost)) if cost is not None else None
    except ArithmeticError:
        cost_value = None
    return CallUsage(
        prompt_tokens=_int_or_none(usage.get("prompt_tokens")),
        completion_tokens=_int_or_none(usage.get("completion_tokens")),
        cached_tokens=_int_or_none(cached) or 0,
        cost=cost_value,
    )


def usage_from_yandex(result: Any) -> CallUsage:
    """Счётчик SDK Yandex AI Studio: `result.usage.input_text_tokens` / `completion_tokens`."""

    usage = getattr(result, "usage", None)
    if usage is None:
        return CallUsage()
    return CallUsage(
        prompt_tokens=_int_or_none(getattr(usage, "input_text_tokens", None)),
        completion_tokens=_int_or_none(getattr(usage, "completion_tokens", None)),
    )


def _int_or_none(value: Any) -> int | None:
    try:
        return int(value) if value is not None else None
    except (TypeError, ValueError):
        return None


def record_call(
    db: Session | None,
    *,
    provider: str,
    model: str | None,
    operation: str,
    prompt_chars: int,
    completion_chars: int,
    usage: CallUsage | None,
    duration_s: float,
    status: str,
    error: str | None = None,
    attempt: int | None = None,
    attempts: int | None = None,
) -> None:
    """Сохраняет одну попытку обращения к модели. Нет счётчика провайдера (сбой до ответа,
    обрыв соединения) — токены оцениваются по длине текста и помечаются как оценка.

    Неудачная попытка дополнительно пишется в журнал ИИ (`ai_log_service`, 14.10.2026):
    сюда сообщают о каждой попытке все клиенты моделей, так что это единственное место, где
    не пропустить ни один сбой. Уровень WARNING — клиент ещё повторит запрос или перейдёт
    на запасную модель; итоговый отказ пишет `ai_client.run_structured` уровнем ERROR.
    Исчерпанный лимит — сразу ERROR: повтор бессмыслен, нужны деньги на ключе."""

    if status != "success":
        log_ai_event(
            db,
            action=STATUS_ACTIONS.get(status, f"Попытка обращения к модели: {status}"),
            error=error,
            level=LogLevel.ERROR if status == "quota" else LogLevel.WARNING,
            result=status,
            provider=provider,
            model=model,
            operation=operation,
            attempt=attempt,
            attempts=attempts,
            duration_s=duration_s,
        )
    if db is None:
        # Клиент вызван без БД (модульные тесты клиентов) — записывать некуда.
        return
    usage = usage or CallUsage()
    estimated = usage.prompt_tokens is None
    prompt_tokens = (
        usage.prompt_tokens if usage.prompt_tokens is not None else round(prompt_chars / CHARS_PER_TOKEN)
    )
    completion_tokens = (
        usage.completion_tokens
        if usage.completion_tokens is not None
        else round(completion_chars / CHARS_PER_TOKEN)
    )
    try:
        with Session(bind=db.get_bind(), join_transaction_mode="create_savepoint") as session:
            job_id = current_job_id()
            job = session.get(BackgroundJob, job_id) if job_id else None
            session.add(
                AiUsageRecord(
                    provider=provider,
                    model=model[:120] if model else None,
                    operation=operation[:80],
                    user_id=current_user_id(),
                    job_id=job_id,
                    job_kind=job.kind if job else None,
                    tender_id=job.tender_id if job else None,
                    prompt_tokens=prompt_tokens,
                    completion_tokens=completion_tokens,
                    cached_tokens=usage.cached_tokens,
                    tokens_estimated=estimated,
                    prompt_chars=prompt_chars,
                    completion_chars=completion_chars,
                    duration_ms=int(duration_s * 1000),
                    status=status,
                    error=error[:MAX_ERROR_CHARS] if error else None,
                    cost=usage.cost,
                )
            )
            session.commit()
    except Exception as exc:  # noqa: BLE001 - учёт не должен ронять обращение к модели
        logger.warning(f"Не удалось записать расход токенов ({provider}, {operation}): {exc}")


# --- детализация ---------------------------------------------------------------------------


def _period_bounds(date_from: date, date_to: date) -> tuple[datetime, datetime]:
    """Границы периода по московскому времени: «сегодня» у пользователя — московские сутки."""

    start = datetime.combine(date_from, datetime.min.time(), tzinfo=MOSCOW_TZ)
    end = datetime.combine(date_to + timedelta(days=1), datetime.min.time(), tzinfo=MOSCOW_TZ)
    return start, end


def _conditions(date_from: date, date_to: date, provider: str | None) -> list:
    start, end = _period_bounds(date_from, date_to)
    conditions = [AiUsageRecord.created_at >= start, AiUsageRecord.created_at < end]
    if provider:
        conditions.append(AiUsageRecord.provider == provider)
    return conditions


def _totals_columns():
    return (
        func.count(AiUsageRecord.id).label("calls"),
        func.coalesce(func.sum(AiUsageRecord.prompt_tokens), 0).label("prompt_tokens"),
        func.coalesce(func.sum(AiUsageRecord.completion_tokens), 0).label("completion_tokens"),
        func.coalesce(func.sum(AiUsageRecord.cached_tokens), 0).label("cached_tokens"),
        func.sum(AiUsageRecord.cost).label("cost"),
        func.coalesce(func.sum(case((AiUsageRecord.status != "success", 1), else_=0)), 0).label("failed"),
        func.coalesce(func.sum(AiUsageRecord.duration_ms), 0).label("duration_ms"),
        # Время и ответ только успешных обращений — для среднего времени ответа и скорости
        # вывода: сбой возвращается за доли секунды без единого токена и тянул бы обе вниз.
        func.coalesce(
            func.sum(case((AiUsageRecord.status == "success", AiUsageRecord.duration_ms), else_=0)), 0
        ).label("success_duration_ms"),
        func.coalesce(
            func.sum(case((AiUsageRecord.status == "success", AiUsageRecord.completion_tokens), else_=0)), 0
        ).label("success_completion_tokens"),
    )


_TOTAL_TOKENS = func.sum(AiUsageRecord.prompt_tokens + AiUsageRecord.completion_tokens)


def _row_totals(row) -> dict:
    prompt = int(row.prompt_tokens or 0)
    completion = int(row.completion_tokens or 0)
    return {
        "calls": int(row.calls or 0),
        "prompt_tokens": prompt,
        "completion_tokens": completion,
        "total_tokens": prompt + completion,
        "cached_tokens": int(row.cached_tokens or 0),
        "cost": float(row.cost) if row.cost is not None else None,
        "failed": int(row.failed or 0),
        "duration_ms": int(row.duration_ms or 0),
        "success_duration_ms": int(row.success_duration_ms or 0),
        "success_completion_tokens": int(row.success_completion_tokens or 0),
    }


def usage_summary(
    db: Session, *, date_from: date, date_to: date, provider: str | None = None
) -> dict:
    """Сводка расхода за период: итог, по дням, по моделям, по шагам разбора, по видам
    задач (со средним на закупку — основа прогноза), по пользователям и самые дорогие
    закупки."""

    conditions = _conditions(date_from, date_to, provider)

    totals = _row_totals(db.execute(select(*_totals_columns()).where(*conditions)).one())
    totals["estimated_calls"] = int(
        db.scalar(
            select(func.count(AiUsageRecord.id)).where(*conditions, AiUsageRecord.tokens_estimated)
        )
        or 0
    )
    totals["tenders"] = int(
        db.scalar(
            select(func.count(func.distinct(AiUsageRecord.tender_id))).where(
                *conditions, AiUsageRecord.tender_id.is_not(None)
            )
        )
        or 0
    )

    local_day = cast(func.timezone("Europe/Moscow", AiUsageRecord.created_at), Date)
    by_day = [
        {"day": row.day.isoformat(), **_row_totals(row)}
        for row in db.execute(
            select(local_day.label("day"), *_totals_columns())
            .where(*conditions)
            .group_by(local_day)
            .order_by(local_day)
        )
    ]

    # Символов на токен — по ответам с настоящим счётчиком провайдера: показывает, насколько
    # консервативна оценка `CHARS_PER_TOKEN`, от которой считается объём текста в запросе.
    # Самый большой запрос — сколько окна реально нужно.
    measured = AiUsageRecord.tokens_estimated.is_(False)
    by_model = []
    for row in db.execute(
        select(
            AiUsageRecord.provider,
            AiUsageRecord.model,
            *_totals_columns(),
            func.sum(case((measured, AiUsageRecord.prompt_chars), else_=0)).label("measured_chars"),
            func.sum(case((measured, AiUsageRecord.prompt_tokens), else_=0)).label("measured_tokens"),
            func.max(AiUsageRecord.prompt_tokens).label("max_prompt_tokens"),
        )
        .where(*conditions)
        .group_by(AiUsageRecord.provider, AiUsageRecord.model)
        .order_by(_TOTAL_TOKENS.desc())
    ):
        measured_tokens = int(row.measured_tokens or 0)
        by_model.append(
            {
                "provider": row.provider,
                "provider_label": PROVIDER_LABELS.get(row.provider, row.provider),
                "model": row.model,
                "chars_per_token": (
                    round(int(row.measured_chars or 0) / measured_tokens, 2) if measured_tokens else None
                ),
                "max_prompt_tokens": int(row.max_prompt_tokens or 0),
                **_row_totals(row),
            }
        )

    by_operation = _grouped_by_label(
        db,
        conditions,
        AiUsageRecord.operation,
        lambda key: OPERATION_LABELS.get(key, key),
        with_prompt_stats=True,
    )
    by_job_kind = _grouped_by_label(
        db,
        conditions,
        AiUsageRecord.job_kind,
        lambda key: JOB_KIND_LABELS.get(key, key) if key else NO_JOB_LABEL,
        with_tenders=True,
    )

    user_rows = db.execute(
        select(AiUsageRecord.user_id, *_totals_columns())
        .where(*conditions)
        .group_by(AiUsageRecord.user_id)
        .order_by(_TOTAL_TOKENS.desc())
    ).all()
    user_names = _user_names(db, {row.user_id for row in user_rows if row.user_id})
    by_user = [
        {
            "key": str(row.user_id) if row.user_id else SYSTEM_USER_KEY,
            "label": (
                user_names.get(row.user_id, "Удалённый пользователь")
                if row.user_id
                else "Система (расписание)"
            ),
            **_row_totals(row),
        }
        for row in user_rows
    ]

    tender_rows = db.execute(
        select(AiUsageRecord.tender_id, *_totals_columns())
        .where(*conditions, AiUsageRecord.tender_id.is_not(None))
        .group_by(AiUsageRecord.tender_id)
        .order_by(_TOTAL_TOKENS.desc())
        .limit(15)
    ).all()
    tenders = _tenders(db, {row.tender_id for row in tender_rows})
    top_tenders = [
        {
            "tender_id": str(row.tender_id),
            "external_id": tenders[row.tender_id].external_id if row.tender_id in tenders else None,
            "title": tenders[row.tender_id].title if row.tender_id in tenders else "Закупка удалена",
            **_row_totals(row),
        }
        for row in tender_rows
    ]

    return {
        "date_from": date_from.isoformat(),
        "date_to": date_to.isoformat(),
        "provider": provider,
        "chars_per_token_estimate": CHARS_PER_TOKEN,
        "totals": totals,
        "by_day": by_day,
        "by_model": by_model,
        "by_operation": by_operation,
        "by_job_kind": by_job_kind,
        "by_user": by_user,
        "top_tenders": top_tenders,
    }


SYSTEM_USER_KEY = "system"
USER_TENDERS_LIMIT = 200


def user_usage(
    db: Session,
    *,
    user_key: str,
    date_from: date,
    date_to: date,
    provider: str | None = None,
) -> dict:
    """Расход одного пользователя (строка «Пользователи» в детализации): на какие шаги и
    разборы ушли его токены и какие закупки он разбирал — со ссылками на них (14.10.2026).

    `user_key` — id пользователя или `system` (расписание, обращения без пользователя), как
    `key` в `by_user` сводки. Закупки — сначала последние разобранные: администратору нужно
    «что он делал», а самые дорогие и так видны в общей сводке. Неверный id — `ValueError`."""

    conditions = _conditions(date_from, date_to, provider)
    if user_key == SYSTEM_USER_KEY:
        conditions.append(AiUsageRecord.user_id.is_(None))
    else:
        conditions.append(AiUsageRecord.user_id == uuid.UUID(user_key))

    totals = _row_totals(db.execute(select(*_totals_columns()).where(*conditions)).one())
    by_operation = _grouped_by_label(
        db, conditions, AiUsageRecord.operation, lambda key: OPERATION_LABELS.get(key, key)
    )
    by_job_kind = _grouped_by_label(
        db,
        conditions,
        AiUsageRecord.job_kind,
        lambda key: JOB_KIND_LABELS.get(key, key) if key else NO_JOB_LABEL,
        with_tenders=True,
    )

    with_tender = [*conditions, AiUsageRecord.tender_id.is_not(None)]
    tenders_total = int(
        db.scalar(select(func.count(func.distinct(AiUsageRecord.tender_id))).where(*with_tender)) or 0
    )
    last_at = func.max(AiUsageRecord.created_at)
    tender_rows = db.execute(
        select(
            AiUsageRecord.tender_id,
            *_totals_columns(),
            last_at.label("last_at"),
            func.array_agg(func.distinct(AiUsageRecord.operation)).label("operations"),
        )
        .where(*with_tender)
        .group_by(AiUsageRecord.tender_id)
        .order_by(last_at.desc())
        .limit(USER_TENDERS_LIMIT)
    ).all()
    tenders = _tenders(db, {row.tender_id for row in tender_rows})
    tender_items = []
    for row in tender_rows:
        tender = tenders.get(row.tender_id)
        tender_items.append(
            {
                "tender_id": str(row.tender_id),
                "external_id": tender.external_id if tender else None,
                "title": tender.title if tender else "Закупка удалена",
                "exists": tender is not None,
                "last_at": row.last_at.isoformat(),
                # Подписи шагов без повторов: четыре схемы AI-оценки — один шаг.
                "operations": sorted({OPERATION_LABELS.get(op, op) for op in row.operations or []}),
                **_row_totals(row),
            }
        )

    return {
        "key": user_key,
        "totals": totals,
        "by_operation": by_operation,
        "by_job_kind": by_job_kind,
        "tenders_total": tenders_total,
        "tenders": tender_items,
    }


def _grouped_by_label(
    db: Session,
    conditions: list,
    column,
    label_of,
    *,
    with_prompt_stats: bool = False,
    with_tenders: bool = False,
) -> list[dict]:
    """Группировка по ключу с человеческой подписью. Разные ключи с одной подписью (четыре
    схемы AI-оценки) сливаются в одну строку — администратору нужен шаг разбора, а не схема.

    Число закупок при слиянии складывается: закупка, по которой шли обе схемы, посчитается
    дважды. Для видов задач (`with_tenders`) это не так — ключ там один на подпись."""

    extra = []
    if with_prompt_stats:
        extra.append(func.max(AiUsageRecord.prompt_tokens).label("max_prompt_tokens"))
    if with_tenders:
        extra.append(func.count(func.distinct(AiUsageRecord.tender_id)).label("tenders"))
    merged: dict[str, dict] = {}
    for row in db.execute(
        select(column.label("key"), *_totals_columns(), *extra).where(*conditions).group_by(column)
    ):
        label = label_of(row.key)
        item = _row_totals(row)
        current = merged.get(label)
        if current is None:
            current = {"key": row.key or "none", "label": label, **item}
            if with_prompt_stats:
                current["max_prompt_tokens"] = 0
            if with_tenders:
                current["tenders"] = 0
            merged[label] = current
        else:
            for name in _SUM_FIELDS:
                current[name] += item[name]
            if item["cost"] is not None:
                current["cost"] = (current["cost"] or 0) + item["cost"]
        if with_prompt_stats:
            current["max_prompt_tokens"] = max(current["max_prompt_tokens"], int(row.max_prompt_tokens or 0))
        if with_tenders:
            current["tenders"] += int(row.tenders or 0)

    rows = sorted(merged.values(), key=lambda item: item["total_tokens"], reverse=True)
    for item in rows:
        if with_prompt_stats:
            item["avg_prompt_tokens"] = round(item["prompt_tokens"] / item["calls"]) if item["calls"] else 0
        if with_tenders:
            # Среднее на закупку — основа прогноза «N закупок в месяц × столько токенов».
            item["avg_tokens_per_tender"] = (
                round(item["total_tokens"] / item["tenders"]) if item["tenders"] else None
            )
    return rows


def _user_names(db: Session, user_ids: set) -> dict:
    if not user_ids:
        return {}
    return {
        user.id: user.full_name or user.username
        for user in db.scalars(select(User).where(User.id.in_(user_ids)))
    }


def _tenders(db: Session, tender_ids: set) -> dict:
    if not tender_ids:
        return {}
    return {tender.id: tender for tender in db.scalars(select(Tender).where(Tender.id.in_(tender_ids)))}


def recent_calls(
    db: Session,
    *,
    date_from: date,
    date_to: date,
    provider: str | None = None,
    limit: int = 50,
    offset: int = 0,
) -> dict:
    """Последние обращения — для разбора «почему этот день дорогой» и сбоев."""

    conditions = _conditions(date_from, date_to, provider)
    total = int(db.scalar(select(func.count(AiUsageRecord.id)).where(*conditions)) or 0)
    records = list(
        db.scalars(
            select(AiUsageRecord)
            .where(*conditions)
            .order_by(AiUsageRecord.created_at.desc())
            .limit(limit)
            .offset(offset)
        )
    )
    tenders = _tenders(db, {record.tender_id for record in records if record.tender_id})
    user_names = _user_names(db, {record.user_id for record in records if record.user_id})
    items = []
    for record in records:
        tender = tenders.get(record.tender_id) if record.tender_id else None
        items.append(
            {
                "id": str(record.id),
                "created_at": record.created_at.isoformat(),
                "provider": record.provider,
                "provider_label": PROVIDER_LABELS.get(record.provider, record.provider),
                "model": record.model,
                "operation": record.operation,
                "operation_label": OPERATION_LABELS.get(record.operation, record.operation),
                "job_kind_label": JOB_KIND_LABELS.get(record.job_kind) if record.job_kind else None,
                "tender_id": str(record.tender_id) if record.tender_id else None,
                "tender_external_id": tender.external_id if tender else None,
                "user": user_names.get(record.user_id) if record.user_id else None,
                "prompt_tokens": record.prompt_tokens,
                "completion_tokens": record.completion_tokens,
                "cached_tokens": record.cached_tokens,
                "tokens_estimated": record.tokens_estimated,
                "prompt_chars": record.prompt_chars,
                "duration_ms": record.duration_ms,
                "status": record.status,
                "status_label": STATUS_LABELS.get(record.status, record.status),
                "error": record.error,
                "cost": float(record.cost) if record.cost is not None else None,
            }
        )
    return {"total": total, "items": items}


def default_period(days: int = 30) -> tuple[date, date]:
    today = datetime.now(MOSCOW_TZ).date()
    return today - timedelta(days=days - 1), today



# Выгрузка «Детализации ИИ» в Excel (15.10.2026). Руководству нужен расход в таблице — сводить
# в своих отчётах и сравнивать модели по скорости и кэшу. Сначала был один CSV с семью
# таблицами подряд — по нему было не сориентироваться, поэтому теперь лист на раздел: шапка
# закреплена, фильтры включены, числа — числами. Обращений за год — десятки тысяч; потолок
# защищает сервер от огромного файла, а упершуюся выгрузку видно на листе «Итог».
EXPORT_CALLS_LIMIT = 50_000
EXPORT_TENDERS_LIMIT = 5_000

_INT = "#,##0"
_DEC1 = "#,##0.0"
_DEC2 = "#,##0.00"
_PCT = "0.0%"
_RUB = "#,##0.00 ₽"
_DATE = "DD.MM.YYYY"
_DATETIME = "DD.MM.YYYY HH:MM:SS"

_HEADER_FILL = PatternFill("solid", fgColor="FF1F2937")
_HEADER_FONT = Font(bold=True, color="FFFFFFFF")
_BOLD = Font(bold=True)
_MUTED = Font(color="FF6B7280")

# Метрики, общие для листов-срезов: (заголовок, формат, значение по строке сводки).
_METRICS: list[tuple[str, str, Any]] = [
    ("Обращений", _INT, lambda r: r["calls"]),
    ("Сбоев", _INT, lambda r: r["failed"]),
    ("Доля сбоев", _PCT, lambda r: r["failed"] / r["calls"] if r["calls"] else None),
    ("Входящие токены", _INT, lambda r: r["prompt_tokens"]),
    ("Исходящие токены", _INT, lambda r: r["completion_tokens"]),
    ("Всего токенов", _INT, lambda r: r["total_tokens"]),
    ("Из кэша", _INT, lambda r: r["cached_tokens"]),
    ("Доля кэша", _PCT, lambda r: r["cached_tokens"] / r["prompt_tokens"] if r["prompt_tokens"] else None),
    (
        "Ср. время ответа, с",
        _DEC2,
        lambda r: (r["success_duration_ms"] / 1000 / (r["calls"] - r["failed"]))
        if r["calls"] > r["failed"]
        else None,
    ),
    ("Скорость вывода, ток/с", _DEC1, lambda r: _speed(r["success_completion_tokens"], r["success_duration_ms"])),
    ("Стоимость", _RUB, lambda r: r["cost"]),
]

_NOTES = [
    ("Входящие токены", "запрос к модели: промпт, документы, карточка закупки"),
    ("Исходящие токены", "ответ модели"),
    ("Из кэша", "входящие токены, которые провайдер взял из кэша (у Claude и DeepSeek в разы дешевле)"),
    ("Доля кэша", "из кэша ÷ входящие"),
    ("Ср. время ответа", "по успешным обращениям: сбой возвращается мгновенно и исказил бы среднее"),
    ("Скорость вывода", "исходящие токены ÷ время успешных ответов; время включает чтение запроса"),
    ("Сбоев", "ошибки сети, лимиты, ответы не по схеме — они тоже расходуют токены"),
    ("Стоимость", "только если её вернул провайдер (шлюз RouterAI), иначе пусто"),
]


def _speed(completion_tokens: int, duration_ms: int) -> float | None:
    return completion_tokens / (duration_ms / 1000) if duration_ms > 0 else None


def _table_sheet(
    workbook: Workbook,
    title: str,
    columns: list[tuple[str, str | None, int]],
    rows: list[list[Any]],
) -> None:
    """Лист-таблица: шапка в первой строке, закреплены шапка и первая колонка, фильтры.
    `columns` — (заголовок, числовой формат или None, ширина)."""

    sheet = workbook.create_sheet(title)
    sheet.append([name for name, _, _ in columns])
    for cell in sheet[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 32
    for row in rows:
        sheet.append(row)
    for index, (_, number_format, width) in enumerate(columns, start=1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = width
        if number_format:
            for (cell,) in sheet.iter_rows(min_row=2, min_col=index, max_col=index):
                cell.number_format = number_format
    sheet.freeze_panes = "B2"
    if rows:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"


def _slice_sheet(workbook, title, head, items, extra=()) -> None:
    """Лист среза сводки: свои колонки слева (`head` — (заголовок, ширина, значение)),
    общие метрики, затем `extra` — (заголовок, формат, значение) только этого среза."""

    columns = [(name, None, width) for name, width, _ in head]
    columns += [(name, fmt, 14) for name, fmt, _ in _METRICS]
    columns += [(name, fmt, 14) for name, fmt, _ in extra]
    rows = [
        [value(item) for _, _, value in head]
        + [value(item) for _, _, value in _METRICS]
        + [value(item) for _, _, value in extra]
        for item in items
    ]
    _table_sheet(workbook, title, columns, rows)


def _naive_msk(moment: datetime) -> datetime:
    # Excel не знает часовых поясов: время пишем московским, как его видит пользователь.
    return moment.astimezone(MOSCOW_TZ).replace(tzinfo=None)


def build_usage_xlsx(
    db: Session, *, date_from: date, date_to: date, provider: str | None = None
) -> tuple[bytes, str, int]:
    """Excel со всей детализацией за период: «Итог» (период, итоговые числа, как читать),
    затем по листу на срез — дни, модели, шаги разбора, виды разбора, пользователи,
    закупки — и «Обращения», по строке на каждое. У каждой строки — входящие и исходящие
    токены, кэш, время ответа и скорость вывода. Возвращает содержимое, имя файла и число
    обращений в файле."""

    summary = usage_summary(db, date_from=date_from, date_to=date_to, provider=provider)
    totals = summary["totals"]
    conditions = _conditions(date_from, date_to, provider)

    records = list(
        db.scalars(
            select(AiUsageRecord)
            .where(*conditions)
            .order_by(AiUsageRecord.created_at.desc())
            .limit(EXPORT_CALLS_LIMIT)
        )
    )
    # Все разобранные закупки, а не 15 самых дорогих, как на экране: в файле их можно
    # отфильтровать и отсортировать самому.
    tender_rows = db.execute(
        select(
            AiUsageRecord.tender_id,
            *_totals_columns(),
            func.max(AiUsageRecord.created_at).label("last_at"),
        )
        .where(*conditions, AiUsageRecord.tender_id.is_not(None))
        .group_by(AiUsageRecord.tender_id)
        .order_by(_TOTAL_TOKENS.desc())
        .limit(EXPORT_TENDERS_LIMIT)
    ).all()
    tenders = _tenders(
        db,
        {row.tender_id for row in tender_rows} | {r.tender_id for r in records if r.tender_id},
    )
    user_names = _user_names(db, {record.user_id for record in records if record.user_id})

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Итог"
    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 26
    sheet.column_dimensions["C"].width = 80
    sheet.append(["Детализация ИИ"])
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.append([])
    info: list[tuple[str, Any, str | None]] = [
        ("Период (МСК)", f"{date_from:%d.%m.%Y} — {date_to:%d.%m.%Y}", None),
        ("Модели", PROVIDER_LABELS.get(provider, provider) if provider else "все", None),
        ("Сформировано", f"{datetime.now(MOSCOW_TZ):%d.%m.%Y %H:%M} МСК", None),
        ("Закупок разобрано", totals["tenders"], _INT),
    ]
    info += [(name, value(totals), fmt) for name, fmt, value in _METRICS]
    info.append(("Токены оценены по длине текста", totals["estimated_calls"], _INT))
    info.append(
        (
            "Обращений на листе «Обращения»",
            len(records),
            _INT,
        )
    )
    for name, value, fmt in info:
        sheet.append([name, value])
        sheet.cell(sheet.max_row, 1).font = _BOLD
        cell = sheet.cell(sheet.max_row, 2)
        cell.alignment = Alignment(horizontal="right")
        if fmt:
            cell.number_format = fmt
    if len(records) < totals["calls"]:
        sheet.cell(sheet.max_row, 3, f"из {totals['calls']:,} — в файл попали последние".replace(",", " "))
    sheet.append([])
    sheet.append(["Как читать"])
    sheet.cell(sheet.max_row, 1).font = _BOLD
    for name, text in _NOTES:
        sheet.append([name, None, text])
        sheet.cell(sheet.max_row, 1).font = _MUTED
        sheet.cell(sheet.max_row, 3).font = _MUTED
    sheet.append([])
    sheet.append(["Листы: По дням · Модели · Шаги разбора · Виды разбора · Пользователи · Закупки · Обращения"])
    sheet.cell(sheet.max_row, 1).font = _MUTED

    _slice_sheet(
        workbook,
        "По дням",
        [("День", 12, lambda r: date.fromisoformat(r["day"]))],
        summary["by_day"],
    )
    for (cell,) in workbook["По дням"].iter_rows(min_row=2, max_col=1):
        cell.number_format = _DATE
    _slice_sheet(
        workbook,
        "Модели",
        [("Провайдер", 14, lambda r: r["provider_label"]), ("Модель", 34, lambda r: r["model"] or "")],
        summary["by_model"],
        [
            ("Символов на токен", _DEC2, lambda r: r["chars_per_token"]),
            ("Макс. запрос, токенов", _INT, lambda r: r["max_prompt_tokens"]),
        ],
    )
    _slice_sheet(
        workbook,
        "Шаги разбора",
        [("Шаг", 30, lambda r: r["label"])],
        summary["by_operation"],
        [
            ("Ср. запрос, токенов", _INT, lambda r: r["avg_prompt_tokens"]),
            ("Макс. запрос, токенов", _INT, lambda r: r["max_prompt_tokens"]),
        ],
    )
    _slice_sheet(
        workbook,
        "Виды разбора",
        [("Вид", 38, lambda r: r["label"])],
        summary["by_job_kind"],
        [
            ("Закупок", _INT, lambda r: r["tenders"]),
            ("Токенов на закупку", _INT, lambda r: r["avg_tokens_per_tender"]),
        ],
    )
    _slice_sheet(workbook, "Пользователи", [("Пользователь", 30, lambda r: r["label"])], summary["by_user"])

    tender_items = []
    for row in tender_rows:
        tender = tenders.get(row.tender_id)
        tender_items.append(
            {
                "external_id": tender.external_id if tender else "",
                "title": tender.title if tender else "Закупка удалена",
                "last_at": _naive_msk(row.last_at),
                **_row_totals(row),
            }
        )
    _slice_sheet(
        workbook,
        "Закупки",
        [
            ("Номер", 22, lambda r: r["external_id"]),
            ("Закупка", 60, lambda r: r["title"]),
            ("Последний разбор", 18, lambda r: r["last_at"]),
        ],
        tender_items,
    )
    for (cell,) in workbook["Закупки"].iter_rows(min_row=2, min_col=3, max_col=3):
        cell.number_format = _DATETIME

    call_rows = []
    for record in records:
        tender = tenders.get(record.tender_id) if record.tender_id else None
        prompt = record.prompt_tokens
        call_rows.append(
            [
                _naive_msk(record.created_at),
                PROVIDER_LABELS.get(record.provider, record.provider),
                record.model or "",
                OPERATION_LABELS.get(record.operation, record.operation),
                JOB_KIND_LABELS.get(record.job_kind, record.job_kind) if record.job_kind else NO_JOB_LABEL,
                tender.external_id if tender else ("удалена" if record.tender_id else ""),
                user_names.get(record.user_id, "Удалённый пользователь") if record.user_id else "Система",
                STATUS_LABELS.get(record.status, record.status),
                prompt,
                record.completion_tokens,
                prompt + record.completion_tokens,
                record.cached_tokens,
                record.cached_tokens / prompt if prompt else None,
                record.duration_ms / 1000,
                _speed(record.completion_tokens, record.duration_ms) if record.status == "success" else None,
                "да" if record.tokens_estimated else "",
                float(record.cost) if record.cost is not None else None,
                record.error or "",
            ]
        )
    _table_sheet(
        workbook,
        "Обращения",
        [
            ("Время (МСК)", _DATETIME, 19),
            ("Провайдер", None, 12),
            ("Модель", None, 30),
            ("Шаг", None, 28),
            ("Разбор", None, 26),
            ("Закупка", None, 22),
            ("Пользователь", None, 22),
            ("Статус", None, 16),
            ("Входящие токены", _INT, 12),
            ("Исходящие токены", _INT, 12),
            ("Всего токенов", _INT, 12),
            ("Из кэша", _INT, 11),
            ("Доля кэша", _PCT, 10),
            ("Время, с", _DEC2, 10),
            ("Скорость вывода, ток/с", _DEC1, 12),
            ("Токены оценены", None, 10),
            ("Стоимость", "#,##0.0000 ₽", 12),
            ("Ошибка", None, 60),
        ],
        call_rows,
    )

    buffer = io.BytesIO()
    workbook.save(buffer)
    file_name = f"sova_ai_usage_{date_from:%Y-%m-%d}_{date_to:%Y-%m-%d}.xlsx"
    return buffer.getvalue(), file_name, len(records)
