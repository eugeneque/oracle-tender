"""Выгрузка тендеров в Excel (раздел 5.7 ТЗ).

Один лист со сводной информацией по выбранному периоду/фильтру, 19 обязательных полей
Приложения D, автофильтр в заголовках, закреплённая шапка, условное форматирование по сроку
подачи и по проценту победителя, гиперссылки на закупку.

Почему `openpyxl` в режиме обычной (не write-only) книги: условное форматирование,
`freeze_panes` и `auto_filter` в write-only режиме недоступны, а объём выгрузки — тысячи
строк, а не миллионы; на 5401 тендере файл собирается за секунды и укладывается в
требование «не более 30 секунд» (REQ-5.7-12) с большим запасом.

Пороги цвета для процента победителя намеренно совпадают с интерфейсом (80/50 — раздел 5.6
ТЗ): отчёт и экран должны красить одну и ту же цифру одинаково, иначе выгрузку перестают
считать достоверной.
"""

from __future__ import annotations

import io
from datetime import date, datetime, timezone
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.ai_profile import AiProfileScore
from app.models.analysis import WinPercentage
from app.models.log import LogLevel
from app.models.manufacturer import Manufacturer
from app.models.region import FederalDistrict, Region, RegionResponsible
from app.models.tender import Tender
from app.models.user import User
from app.core.timezones import now_msk
from app.services.audit import log_action
from app.services.tender_service import TenderFilters, apply_tender_filters

# Порядок колонок — ровно Приложение D ТЗ, 19 полей. Менять порядок нельзя: заказчик
# сверяет выгрузку с приложением по позициям, а не по названиям.
# (заголовок, ширина колонки)
COLUMNS: list[tuple[str, int]] = [
    ("Статус", 16),
    ("Начало приёма заявок", 20),
    ("Окончание приёма заявок", 22),
    ("Осталось дней", 14),
    ("Месяц", 10),
    ("Сумма", 18),
    ("Валюта", 9),
    ("Тип конкурса", 22),
    ("Способ закупки", 24),
    ("Организатор", 40),
    ("ФО", 26),
    ("Регион организатора", 22),
    ("Регион поставки", 22),
    ("Ответственный за регион", 24),
    ("Руководитель ответственный за регион", 30),
    ("Наименование ЭТП", 20),
    ("Ссылка", 34),
    ("ID сделки", 20),
    ("Комментарий", 60),
]

# Дополнительные колонки, идущие ПОСЛЕ обязательных девятнадцати (Приложение D ТЗ,
# добавлены 03.09.2026). Порядок обязательных при этом не меняется — заказчик сверяет
# выгрузку с приложением по позициям.
EXTRA_COLUMNS: list[tuple[str, int]] = [
    ("AI-оценка", 12),
    ("Вердикт", 22),
    ("Этап", 18),
    ("Ответственный за тендер", 24),
    # Процент соответствия характеристик (раздел 5.5.2) — не главная метрика с 03.09.2026,
    # но из выгрузки не убирается: по нему в отделе сравнивают приборы.
    ("% соответствия МИРТЕК", 18),
]

# Номер колонки итоговой AI-оценки — по ней идёт условное форматирование (раздел 5.7 ТЗ).
AI_SCORE_COLUMN = len(COLUMNS) + 1
WIN_PERCENTAGE_COLUMN = len(COLUMNS) + len(EXTRA_COLUMNS)

# Строка, с которой начинаются данные: 1-3 — шапка файла (название, дата, период),
# 4 — пустая, 5 — заголовки таблицы, 6 — первая запись.
HEADER_ROW = 5
FIRST_DATA_ROW = HEADER_ROW + 1

VERDICT_LABELS_XLSX = {
    "go": "ИДТИ",
    "go_with_reservations": "ИДТИ С ОГОВОРКАМИ",
    "no_go": "НЕ ИДТИ",
}
STAGE_LABELS_XLSX = {
    "ai_selected": "Новая",
    "under_review": "На проверке",
    "application_submitted": "Заявка подана",
    "won": "Выиграли",
    "lost": "Проиграли",
    "rejected": "Отклонён",
}
STATUS_LABELS = {
    "collecting_bids": "Сбор заявок",
    "evaluation": "Оценка",
    "completed": "Завершено",
    "cancelled": "Отменено",
}
TENDER_TYPE_LABELS = {
    "supply_only": "Поставка ИПУ",
    "complex": "Комплекс (ПУ + работы)",
    "works_only": "Работы (СМР и ПНР)",
    "reverification": "Переповерка",
    "other": "Прочее",
}
MONTH_NAMES = (
    "январь", "февраль", "март", "апрель", "май", "июнь",
    "июль", "август", "сентябрь", "октябрь", "ноябрь", "декабрь",
)

_HEADER_FILL = PatternFill("solid", fgColor="FF1F2937")
_TITLE_FONT = Font(bold=True, size=14)
_HEADER_FONT = Font(bold=True, color="FFFFFFFF")
_THIN = Side(style="thin", color="FFD9D9D9")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

# Заливки условного форматирования — те же смыслы, что и цвета в интерфейсе.
_FILL_GREEN = PatternFill("solid", fgColor="FFC6EFCE")
_FILL_YELLOW = PatternFill("solid", fgColor="FFFFEB9C")
_FILL_RED = PatternFill("solid", fgColor="FFFFC7CE")
_FONT_GREEN = Font(color="FF006100")
_FONT_YELLOW = Font(color="FF9C6500")
_FONT_RED = Font(color="FF9C0006")


def _days_left(application_end: datetime | None) -> int | None:
    """Дней до окончания приёма заявок по календарным датам.

    Сравниваются именно даты, а не полные метки времени: иначе дедлайн «сегодня в 18:00»
    после полуночи показывал бы -1, хотя заявку ещё принимают. Та же логика, что во
    фронтенде (`utils/format.ts::daysLeft`) — расхождение отчёта с экраном недопустимо.
    """

    if application_end is None:
        return None
    today = datetime.now(timezone.utc).date()
    return (application_end.date() - today).days


def _month_label(publish_date: date | None) -> str | None:
    if publish_date is None:
        return None
    return MONTH_NAMES[publish_date.month - 1]


def _period_label(filters: TenderFilters) -> str:
    """Человекочитаемый период отбора для шапки файла (REQ-5.7-10).

    Приоритет у даты публикации: именно она отвечает на вопрос «за какой период отчёт».
    Если период не задан ни одним фильтром — так и пишем, чтобы никто не принял выгрузку
    за отчёт «за месяц».
    """

    def _fmt(value: date | None) -> str:
        return value.strftime("%d.%m.%Y") if value else "…"

    if filters.publish_date_from or filters.publish_date_to:
        return f"дата размещения с {_fmt(filters.publish_date_from)} по {_fmt(filters.publish_date_to)}"
    if filters.deadline_from or filters.deadline_to:
        return f"окончание приёма заявок с {_fmt(filters.deadline_from)} по {_fmt(filters.deadline_to)}"
    return "весь период (ограничение по датам не задано)"


def _load_reference_data(db: Session) -> tuple[dict, dict, dict]:
    """Справочники одним заходом: регионы, федеральные округа, ответственные.

    Тендеров в выгрузке тысячи, а регионов 89 — тянуть справочник на каждую строку значило
    бы получить тысячи лишних запросов.
    """

    regions = {row.code: row.name for row in db.scalars(select(Region))}
    districts = {row.code: row.name for row in db.scalars(select(FederalDistrict))}
    region_district = {row.code: row.federal_district_code for row in db.scalars(select(Region))}
    responsibles = {
        row.region_code: (row.responsible_name, row.manager_name)
        for row in db.scalars(select(RegionResponsible))
    }
    districts_by_region = {
        code: districts.get(district_code) for code, district_code in region_district.items()
    }
    return regions, districts_by_region, responsibles


def _tender_row(
    tender: Tender,
    *,
    regions: dict,
    districts_by_region: dict,
    responsibles: dict,
) -> list:
    """Одна строка выгрузки в порядке Приложения D."""

    organizer_code = tender.region_organizer_code
    # Регион поставки: если не заполнен — дублируется регион организатора (Приложение D
    # и раздел 5.4 п.6 ТЗ). Подстановка делается здесь, а не в БД: пустое поле в базе
    # означает «не определено», и затирать это значение молчаливым дублем нельзя.
    delivery_code = tender.region_delivery_code or organizer_code
    responsible, manager = responsibles.get(organizer_code, (None, None))

    return [
        STATUS_LABELS.get(tender.status or "", tender.status),
        tender.application_start.replace(tzinfo=None) if tender.application_start else None,
        tender.application_end.replace(tzinfo=None) if tender.application_end else None,
        _days_left(tender.application_end),
        _month_label(tender.publish_date),
        float(tender.price) if tender.price is not None else None,
        tender.currency,
        TENDER_TYPE_LABELS.get(tender.tender_type or "", tender.tender_type),
        tender.procurement_method,
        tender.organizer_name or tender.customer_name,
        districts_by_region.get(organizer_code),
        regions.get(organizer_code),
        regions.get(delivery_code),
        responsible,
        manager,
        tender.source.name if tender.source else None,
        tender.source_url,
        tender.external_id,
        tender.ai_comment,
    ]


def _ai_scores(db: Session, tender_ids: list) -> dict:
    """Текущие AI-оценки по тендерам — одним запросом на всю выгрузку, как и проценты."""

    if not tender_ids:
        return {}
    rows = db.execute(
        select(AiProfileScore.tender_id, AiProfileScore.overall_score, AiProfileScore.verdict)
        .where(
            AiProfileScore.tender_id.in_(tender_ids), AiProfileScore.is_current.is_(True)
        )
    ).all()
    return {tender_id: (overall, verdict) for tender_id, overall, verdict in rows}


def _assignee_names(db: Session, tenders: list[Tender]) -> dict:
    ids = [tender.assignee_id for tender in tenders if tender.assignee_id]
    if not ids:
        return {}
    return dict(db.execute(select(User.id, User.full_name).where(User.id.in_(ids))).all())


def _win_percentages(db: Session, tender_ids: list) -> dict:
    """Процент победителя МИРТЕК по тендерам — одним запросом на всю выгрузку.

    В Приложении D этой колонки нет, но условное форматирование по проценту победителя
    требование 5.7-07 предписывает, а красить нечего, если числа в файле не будет. Поэтому
    колонка добавляется ПОСЛЕ девятнадцати обязательных — порядок Приложения D не ломается.
    """

    if not tender_ids:
        return {}
    rows = db.execute(
        select(WinPercentage.tender_id, WinPercentage.percentage)
        .join(Manufacturer, Manufacturer.id == WinPercentage.manufacturer_id)
        .where(
            WinPercentage.tender_id.in_(tender_ids),
            Manufacturer.is_mirtek.is_(True),
            WinPercentage.is_current.is_(True),
        )
    ).all()
    return {tender_id: percentage for tender_id, percentage in rows}


def build_tenders_workbook(
    db: Session, filters: TenderFilters, *, limit: int | None = None
) -> tuple[bytes, int]:
    """Собирает XLSX и возвращает (содержимое файла, число выгруженных строк)."""

    query = apply_tender_filters(select(Tender), filters).order_by(
        Tender.publish_date.desc().nullslast(), Tender.created_at.desc()
    )
    if limit is not None:
        query = query.limit(limit)
    tenders = list(db.scalars(query))

    regions, districts_by_region, responsibles = _load_reference_data(db)
    tender_ids = [tender.id for tender in tenders]
    percentages = _win_percentages(db, tender_ids)
    scores = _ai_scores(db, tender_ids)
    assignees = _assignee_names(db, tenders)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Тендеры"

    # --- Шапка файла: название, дата формирования, период отбора (REQ-5.7-09, 5.7-10).
    sheet.cell(row=1, column=1, value="САСТ-Тендеры — сводная выгрузка").font = _TITLE_FONT
    sheet.cell(row=2, column=1, value="Дата формирования:")
    # МСК явно, а не по часам сервера: в Docker он обычно живёт в UTC, и дата
    # формирования отчёта разошлась бы с журналом на три часа (раздел 8 ТЗ).
    sheet.cell(row=2, column=2, value=now_msk().strftime("%d.%m.%Y %H:%M"))
    sheet.cell(row=3, column=1, value="Период отбора:")
    sheet.cell(row=3, column=2, value=_period_label(filters))
    for row in (2, 3):
        sheet.cell(row=row, column=1).font = Font(bold=True)

    # --- Заголовки таблицы.
    headers = [title for title, _ in COLUMNS] + [title for title, _ in EXTRA_COLUMNS]
    for index, title in enumerate(headers, start=1):
        cell = sheet.cell(row=HEADER_ROW, column=index, value=title)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = _BORDER

    # --- Данные.
    for offset, tender in enumerate(tenders):
        row_number = FIRST_DATA_ROW + offset
        values = _tender_row(
            tender,
            regions=regions,
            districts_by_region=districts_by_region,
            responsibles=responsibles,
        )
        overall, verdict = scores.get(tender.id, (None, None))
        values.append(float(overall) if overall is not None else None)
        values.append(VERDICT_LABELS_XLSX.get(verdict or "", verdict))
        values.append(STAGE_LABELS_XLSX.get(tender.stage or "", tender.stage))
        values.append(assignees.get(tender.assignee_id) if tender.assignee_id else None)
        percentage = percentages.get(tender.id)
        values.append(float(percentage) if percentage is not None else None)

        for index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_number, column=index, value=value)
            cell.border = _BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=index in (10, 19))

        # Даты и деньги — форматами ячеек, а не строками: иначе в Excel по ним нельзя ни
        # отсортировать, ни посчитать итог.
        sheet.cell(row=row_number, column=2).number_format = "DD.MM.YYYY HH:MM"
        sheet.cell(row=row_number, column=3).number_format = "DD.MM.YYYY HH:MM"
        sheet.cell(row=row_number, column=6).number_format = "# ##0.00"
        sheet.cell(row=row_number, column=AI_SCORE_COLUMN).number_format = "0.0"
        sheet.cell(row=row_number, column=WIN_PERCENTAGE_COLUMN).number_format = "0.0"

        # Гиперссылка на закупку (REQ-5.7-08) — на самой ячейке со ссылкой.
        if tender.source_url:
            link_cell = sheet.cell(row=row_number, column=17)
            link_cell.hyperlink = tender.source_url
            link_cell.value = "Открыть закупку"
            link_cell.font = Font(color="FF0563C1", underline="single")

    last_row = max(FIRST_DATA_ROW + len(tenders) - 1, FIRST_DATA_ROW)

    # --- Ширины колонок и высота строки заголовков.
    for index, (_, width) in enumerate(COLUMNS + EXTRA_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[HEADER_ROW].height = 34

    # --- Автофильтр в заголовках (REQ-5.7-04) и закрепление шапки (REQ-5.7-05).
    last_column = get_column_letter(len(headers))
    sheet.auto_filter.ref = f"A{HEADER_ROW}:{last_column}{last_row}"
    sheet.freeze_panes = f"A{FIRST_DATA_ROW}"

    if tenders:
        _apply_conditional_formatting(sheet, last_row=last_row)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue(), len(tenders)


def _apply_conditional_formatting(sheet, *, last_row: int) -> None:
    """Условное форматирование: срок подачи, AI-оценка и процент соответствия
    (REQ-5.7-06, 5.7-07; раздел 5.7 ТЗ в редакции 03.09.2026 — цвет по `overall_score`).

    Правила именно условные, а не «покрасить ячейку при генерации»: пользователь правит
    выгрузку руками и пересортировывает её, и заливка должна следовать за значением, а не
    оставаться на месте.
    """

    # Осталось дней: меньше 5 — красный, 5-14 — жёлтый. Отрицательные значения (срок вышел)
    # попадают в первое правило и тоже краснеют, что верно по смыслу.
    days_range = f"D{FIRST_DATA_ROW}:D{last_row}"
    sheet.conditional_formatting.add(
        days_range,
        CellIsRule(operator="lessThan", formula=["5"], fill=_FILL_RED, font=_FONT_RED),
    )
    sheet.conditional_formatting.add(
        days_range,
        CellIsRule(
            operator="between", formula=["5", "14"], fill=_FILL_YELLOW, font=_FONT_YELLOW
        ),
    )

    # Пороги те же, что в интерфейсе (раздел 5.6 ТЗ): зелёный ≥80, жёлтый 50-80, красный
    # ниже. Красятся обе процентные колонки — главная метрика и процент соответствия.
    for column_index in (AI_SCORE_COLUMN, WIN_PERCENTAGE_COLUMN):
        column = get_column_letter(column_index)
        percent_range = f"{column}{FIRST_DATA_ROW}:{column}{last_row}"
        sheet.conditional_formatting.add(
            percent_range,
            CellIsRule(
                operator="greaterThanOrEqual", formula=["80"], fill=_FILL_GREEN, font=_FONT_GREEN
            ),
        )
        sheet.conditional_formatting.add(
            percent_range,
            CellIsRule(
                operator="between", formula=["50", "79.999"], fill=_FILL_YELLOW, font=_FONT_YELLOW
            ),
        )
        sheet.conditional_formatting.add(
            percent_range,
            CellIsRule(operator="lessThan", formula=["50"], fill=_FILL_RED, font=_FONT_RED),
        )


def export_tenders(
    db: Session, filters: TenderFilters, actor: User, *, limit: int | None = None
) -> tuple[bytes, str, int]:
    """Выгрузка с журналированием (REQ-5.9-07). Возвращает (файл, имя файла, число строк)."""

    started = datetime.now(timezone.utc)
    content, rows = build_tenders_workbook(db, filters, limit=limit)
    elapsed = (datetime.now(timezone.utc) - started).total_seconds()

    file_name = f"sast-tenders-{now_msk().strftime('%Y%m%d-%H%M')}.xlsx"
    log_action(
        db,
        component="export",
        action="export_tenders_xlsx",
        result="success",
        level=LogLevel.INFO,
        details=(
            f"Строк выгружено: {rows}; период: {_period_label(filters)}; "
            f"время формирования: {elapsed:.1f} с; файл: {file_name}"
        ),
        user_id=actor.id,
    )
    db.commit()
    return content, file_name, rows
