"""Отчёт «Рынок ПУ» по поверкам ФГИС «Аршин» (раздел 79 ARCHITECTURE.md, раздел 8 ТЗ).

Четыре экрана повторяют страницы Power BI аналитика: общая информация и интерфейсы,
исполнения по годам, группы компаний по годам, локализация по кварталам. Все цифры — из
помесячного агрегата `meter_verification_monthly` (пригодные неаннулированные поверки по
варианту написания), соединённого со справочником: производитель — владелец описания типа
(или переопределение на модификации), атрибуты — у модификации. Только типы «в отчёте»
(правило 6 ТЗ).

Агрегат за период — несколько тысяч строк «производитель × исполнение × 1ф/3ф × интерфейс»,
поэтому перекрёстные фильтры (клик по производителю сужает остальные блоки, как в Power BI)
считаются в Python над этими строками, а не отдельным запросом на каждый блок.

Неразмеченные варианты не пропадают: их поверки видны строкой «не размечено» (раздел 6.3 ТЗ),
а их количество отдаётся отдельно — на экране видно, что итог неполный.
"""

from __future__ import annotations

import io
import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any

from sqlalchemy import text
from sqlalchemy.orm import Session

from app.models.meter_market import DESIGN_LABELS

UNMAPPED = "не размечено"
FIRST_YEAR = 2021

_MANUFACTURER_SQL = "coalesce(m.manufacturer_override_id, t.manufacturer_id)"
_BASE_FROM = (
    "FROM meter_verification_monthly mv "
    "JOIN meter_market_variants v ON v.id = mv.variant_id "
    "JOIN meter_market_types t ON t.id = v.type_id AND t.in_report "
    "LEFT JOIN meter_market_modifications m ON m.id = v.modification_id "
)


# ── Период ─────────────────────────────────────────────────────────────────────────────────


@dataclass(frozen=True)
class Period:
    """Полуинтервал месяцев `[start, end)` и подпись для заголовка."""

    start: date
    end: date
    label: str


def _add_months(value: date, months: int) -> date:
    index = value.year * 12 + value.month - 1 + months
    return date(index // 12, index % 12 + 1, 1)


def quarter_label(year: int, quarter: int) -> str:
    return f"{quarter} кв. {year}"


def data_edge(db: Session) -> date | None:
    """Конец последнего полного квартала по дате последнего загруженного пакета.

    Отчёт аналитика — поквартальный, а изменения ФГИС приходят каждый день: без этой границы
    «2026 год» к концу октября показывал бы три квартала и кусок четвёртого."""

    package = db.execute(
        text("SELECT package_id FROM meter_market_loads WHERE status = 'done' ORDER BY package_id DESC LIMIT 1")
    ).scalar()
    if not package:
        return None
    loaded = datetime.strptime(package, "%Y%m%d").date()
    return date(loaded.year, (loaded.month - 1) // 3 * 3 + 1, 1)


def parse_period(value: str | None, edge: date | None) -> Period:
    """`year:2026`, `ytd:2026`, `quarter:2026-3`, `month:2026-09`, `all`. Без значения — год
    последнего полного квартала с начала года («2026 г., 1–3 кв.»)."""

    edge = edge or date.today().replace(day=1)
    last = _add_months(edge, -1)  # последний полный месяц
    if not value:
        value = f"ytd:{last.year}"
    kind, _, arg = value.partition(":")
    try:
        if kind == "all":
            return Period(date(FIRST_YEAR, 1, 1), edge, f"{FIRST_YEAR} г. – {_edge_label(edge)}")
        if kind in ("year", "ytd"):
            year = int(arg)
            start, end = date(year, 1, 1), date(year + 1, 1, 1)
            if end > edge:
                end = edge
                return Period(start, end, f"{year} г. ({_quarters_label(start, end)})")
            return Period(start, end, f"{year} г.")
        if kind == "quarter":
            year, quarter = (int(x) for x in arg.split("-"))
            start = date(year, (quarter - 1) * 3 + 1, 1)
            return Period(start, _add_months(start, 3), f"{quarter_label(year, quarter)} г.")
        if kind == "month":
            year, month = (int(x) for x in arg.split("-"))
            start = date(year, month, 1)
            return Period(start, _add_months(start, 1), f"{start:%m.%Y}")
    except ValueError as exc:
        raise ValueError(f"Непонятный период «{value}»") from exc
    raise ValueError(f"Непонятный период «{value}»")


def _edge_label(edge: date) -> str:
    last = _add_months(edge, -1)
    quarter = (last.month - 1) // 3 + 1
    return f"{quarter_label(last.year, quarter)} г." if last.month % 3 == 0 else f"{last:%m.%Y}"


def _quarters_label(start: date, end: date) -> str:
    last = _add_months(end, -1)
    q_to = (last.month - 1) // 3 + 1
    return f"1–{q_to} кв." if q_to > 1 else "1 кв."


# ── Справочные данные отчёта ───────────────────────────────────────────────────────────────


def _manufacturer_names(db: Session) -> dict[int, str]:
    return dict(db.execute(text("SELECT id, name FROM meter_market_manufacturers")).all())


def thresholds(db: Session) -> list[tuple[int, float]]:
    return [
        (int(y), float(s))
        for y, s in db.execute(text("SELECT year, min_score FROM meter_market_score_thresholds ORDER BY year")).all()
    ]


def requirement_year(score: float | None, limits: list[tuple[int, float]]) -> int | None:
    """«Соответствие требованиям года» — наибольший год, чей порог не выше баллов (правило 3 ТЗ)."""

    if score is None:
        return None
    met = [year for year, minimum in limits if score >= minimum]
    return max(met) if met else None


def status(db: Session) -> dict[str, Any]:
    last = db.execute(
        text(
            "SELECT package_id, kind, finished_at, message FROM meter_market_loads "
            "WHERE status = 'done' ORDER BY package_id DESC LIMIT 1"
        )
    ).mappings().first()
    running = db.execute(
        text(
            "SELECT package_id, kind, status, checkpoint, started_at, message FROM meter_market_loads "
            "WHERE status <> 'done' ORDER BY id DESC LIMIT 1"
        )
    ).mappings().first()
    edge = data_edge(db)
    return {
        "last_package": dict(last) if last else None,
        "pending_package": dict(running) if running else None,
        "edge": edge,
        "edge_label": _edge_label(edge) if edge else None,
        "new_types": db.execute(text("SELECT count(*) FROM meter_market_types WHERE review_status = 'new'")).scalar(),
        "unmapped_variants": db.execute(
            text(
                "SELECT count(*) FROM meter_market_variants v JOIN meter_market_types t ON t.id = v.type_id "
                "WHERE t.in_report AND v.modification_id IS NULL"
            )
        ).scalar(),
        "thresholds": [{"year": y, "min_score": s} for y, s in thresholds(db)],
    }


# ── Экран 1: общая информация и интерфейсы ─────────────────────────────────────────────────


@dataclass(frozen=True)
class Filters:
    manufacturer_id: int | None = None
    design: str | None = None
    phase: str | None = None
    interface: str | None = None


def _cells(db: Session, period: Period, extra_group: str = "", extra_select: str = "") -> list[dict[str, Any]]:
    rows = db.execute(
        text(
            f"SELECT {_MANUFACTURER_SQL} AS manufacturer_id, m.design, m.phase, m.interface, "
            f"(v.modification_id IS NULL) AS unmapped{extra_select}, sum(mv.count) AS count "
            f"{_BASE_FROM} WHERE mv.month >= :start AND mv.month < :end "
            f"GROUP BY 1, 2, 3, 4, 5{extra_group}"
        ),
        {"start": period.start, "end": period.end},
    ).mappings()
    return [dict(r) for r in rows]


def _design_label(cell: dict[str, Any]) -> str:
    if cell["unmapped"]:
        return UNMAPPED
    return DESIGN_LABELS.get(cell["design"] or "nd", "н/д")


def _phase_label(cell: dict[str, Any]) -> str:
    return UNMAPPED if cell["unmapped"] else (cell["phase"] or "н/д")


def _interface_label(cell: dict[str, Any]) -> str:
    return UNMAPPED if cell["unmapped"] else (cell["interface"] or "н/д")


def _keep(cell: dict[str, Any], filters: Filters, names: dict[int, str], skip: str) -> bool:
    if skip != "manufacturer" and filters.manufacturer_id is not None and cell["manufacturer_id"] != filters.manufacturer_id:
        return False
    if skip != "design" and filters.design is not None and _design_label(cell) != filters.design:
        return False
    if skip != "phase" and filters.phase is not None and _phase_label(cell) != filters.phase:
        return False
    if skip != "interface" and filters.interface is not None and _interface_label(cell) != filters.interface:
        return False
    return True


def _breakdown(items: dict[Any, int]) -> list[dict[str, Any]]:
    total = sum(items.values())
    return [
        {"key": key, "count": count, "share": round(count * 100 / total, 2) if total else 0.0}
        for key, count in sorted(items.items(), key=lambda kv: -kv[1])
    ]


def overview(db: Session, period: Period, filters: Filters) -> dict[str, Any]:
    names = _manufacturer_names(db)
    cells = _cells(db, period)

    def block(skip: str, label) -> dict[Any, int]:
        acc: dict[Any, int] = defaultdict(int)
        for cell in cells:
            if _keep(cell, filters, names, skip):
                acc[label(cell)] += int(cell["count"])
        return acc

    manufacturers = _breakdown(block("manufacturer", lambda c: c["manufacturer_id"]))
    for row in manufacturers:
        row["name"] = names.get(row["key"], "Производитель не указан")
    total_all = sum(int(c["count"]) for c in cells if _keep(c, filters, names, ""))
    unmapped = sum(int(c["count"]) for c in cells if c["unmapped"] and _keep(c, filters, names, ""))
    return {
        "period": {"start": period.start, "end": period.end, "label": period.label},
        "total": total_all,
        "unmapped": unmapped,
        "manufacturers": manufacturers,
        "designs": _breakdown(block("design", _design_label)),
        "phases": _breakdown(block("phase", _phase_label)),
        "interfaces": _breakdown(block("interface", _interface_label)),
    }


# ── Экраны 2 и 3: по годам ─────────────────────────────────────────────────────────────────


def _years(edge: date) -> list[int]:
    return list(range(FIRST_YEAR, _add_months(edge, -1).year + 1))


def designs_by_year(db: Session, edge: date, phase: str | None = None) -> dict[str, Any]:
    """Производитель → исполнения, столбцы — годы (экран 2). Фильтр — 1ф/3ф."""

    names = _manufacturer_names(db)
    period = Period(date(FIRST_YEAR, 1, 1), edge, f"{FIRST_YEAR} г. – {_edge_label(edge)}")
    cells = _cells(db, period, ", 6", ", extract(year FROM mv.month)::int AS year")
    years = _years(edge)
    table: dict[int, dict[str, dict[int, int]]] = defaultdict(lambda: defaultdict(lambda: defaultdict(int)))
    for cell in cells:
        if phase and _phase_label(cell) != phase:
            continue
        table[cell["manufacturer_id"]][_design_label(cell)][cell["year"]] += int(cell["count"])
    rows = []
    for manufacturer_id, designs in table.items():
        by_year = {y: sum(d.get(y, 0) for d in designs.values()) for y in years}
        rows.append(
            {
                "manufacturer_id": manufacturer_id,
                "name": names.get(manufacturer_id, "Производитель не указан"),
                "years": by_year,
                "total": sum(by_year.values()),
                "designs": sorted(
                    (
                        {"design": design, "years": {y: values.get(y, 0) for y in years}, "total": sum(values.values())}
                        for design, values in designs.items()
                    ),
                    key=lambda r: -r["total"],
                ),
            }
        )
    rows.sort(key=lambda r: -r["total"])
    totals = {y: sum(r["years"][y] for r in rows) for y in years}
    return {"label": period.label, "years": years, "rows": rows, "totals": totals, "total": sum(totals.values())}


def _groups(db: Session) -> tuple[dict[int, str], list[tuple[int, int, int | None, int | None]]]:
    groups = dict(db.execute(text("SELECT id, name FROM meter_market_groups")).all())
    members = [
        tuple(r)
        for r in db.execute(text("SELECT group_id, manufacturer_id, year_from, year_to FROM meter_market_group_members")).all()
    ]
    return groups, members


def group_for(manufacturer_id: int | None, year: int, members: list[tuple[int, int, int | None, int | None]]) -> int | None:
    for group_id, member_id, year_from, year_to in members:
        if member_id == manufacturer_id and (year_from is None or year >= year_from) and (year_to is None or year <= year_to):
            return group_id
    return None


def groups_by_year(db: Session, edge: date) -> dict[str, Any]:
    """Группа → входящие компании, столбцы — годы и «Всего» (экран 3). Состав группы — на год
    поверки: АтомЦифроСбыт в 2022–2025 в «Энергомере», с 2026 — в «Заводе НАРТИС»."""

    names = _manufacturer_names(db)
    groups, members = _groups(db)
    period = Period(date(FIRST_YEAR, 1, 1), edge, f"{FIRST_YEAR} г. – {_edge_label(edge)}")
    years = _years(edge)
    rows = db.execute(
        text(
            f"SELECT {_MANUFACTURER_SQL} AS manufacturer_id, extract(year FROM mv.month)::int AS year, sum(mv.count) "
            f"{_BASE_FROM} WHERE mv.month >= :start AND mv.month < :end GROUP BY 1, 2"
        ),
        {"start": period.start, "end": period.end},
    ).all()
    table: dict[str, dict[str, dict[int, int]]] = defaultdict(lambda: defaultdict(lambda: defaultdict(int)))
    for manufacturer_id, year, count in rows:
        name = names.get(manufacturer_id, "Производитель не указан")
        group_id = group_for(manufacturer_id, year, members)
        group = groups[group_id] if group_id is not None else name
        table[group][name][year] += int(count)
    result = []
    for group, companies in table.items():
        by_year = {y: sum(c.get(y, 0) for c in companies.values()) for y in years}
        result.append(
            {
                "name": group,
                "years": by_year,
                "total": sum(by_year.values()),
                "companies": sorted(
                    (
                        {"name": company, "years": {y: v.get(y, 0) for y in years}, "total": sum(v.values())}
                        for company, v in companies.items()
                    ),
                    key=lambda r: -r["total"],
                ),
            }
        )
    result.sort(key=lambda r: -r["total"])
    totals = {y: sum(r["years"][y] for r in result) for y in years}
    return {"label": period.label, "years": years, "rows": result, "totals": totals, "total": sum(totals.values())}


# ── Экран 4: локализация ───────────────────────────────────────────────────────────────────

LOCALIZATION_FROM_YEAR = 2025


def localization(db: Session, edge: date, requirement: int | None = None) -> dict[str, Any]:
    """Поверенные локализованные ПУ: производитель → баллы, столбцы — кварталы с 2025 года.

    Точка отсчёта — дата внесения исполнения в РПП (правило 4 ТЗ): поверки раньше неё не
    считаются. Если у исполнения есть история записей ГИСП (`meter_market_rpp_records`: 83 →
    95 → 113 баллов), поверка относится к записи, действовавшей на дату поверки, — и к её
    «требованиям года». Без истории — баллы модификации с даты `rpp_included_at`; если и даты
    нет (в файле аналитика её нет), поверки считаются с начала периода, а их количество
    отдаётся отдельно (`undated`) — расхождение видно, а не спрятано. Месяц, внутри которого
    меняется запись, делится точно — по сырым поверкам."""

    names = _manufacturer_names(db)
    limits = thresholds(db)
    start = date(LOCALIZATION_FROM_YEAR, 1, 1)
    quarters: list[str] = []
    cursor = start
    while cursor < edge:
        quarters.append(f"{cursor.year}-{(cursor.month - 1) // 3 + 1}")
        cursor = _add_months(cursor, 3)

    mods = db.execute(
        text(
            f"SELECT m.id, {_MANUFACTURER_SQL} AS manufacturer_id, m.score, m.rpp_included_at "
            "FROM meter_market_modifications m JOIN meter_market_types t ON t.id = m.type_id AND t.in_report "
            "WHERE m.in_rpp"
        )
    ).mappings().all()
    history: dict[int, list[tuple[date, float | None]]] = defaultdict(list)
    for mod_id, issued, score in db.execute(
        text("SELECT modification_id, issued_at, score FROM meter_market_rpp_records ORDER BY issued_at")
    ):
        history[mod_id].append((issued, score))

    # Отрезки «с даты — баллы» по каждому исполнению.
    segments: dict[int, list[tuple[date, float | None]]] = {}
    undated_mods: set[int] = set()
    for mod in mods:
        if history.get(mod["id"]):
            segments[mod["id"]] = history[mod["id"]]
        elif mod["score"] is not None:
            if mod["rpp_included_at"] is None:
                undated_mods.add(mod["id"])
            segments[mod["id"]] = [(mod["rpp_included_at"] or date.min, mod["score"])]
    by_id = {mod["id"]: mod for mod in mods if mod["id"] in segments}
    empty = {"quarters": quarters, "rows": [], "totals": {q: 0 for q in quarters}, "total": 0, "undated": 0,
             "requirement": requirement, "requirement_years": sorted({y for y, _ in limits})}
    if not by_id:
        return empty

    def active(mod_id: int, day: date) -> float | None | bool:
        """Баллы записи, действующей на дату; `False` — исполнение ещё не внесено."""

        current: float | None | bool = False
        for issued, score in segments[mod_id]:
            if issued <= day:
                current = score
        return current

    monthly = db.execute(
        text(
            "SELECT v.modification_id, mv.month, sum(mv.count) FROM meter_verification_monthly mv "
            "JOIN meter_market_variants v ON v.id = mv.variant_id "
            "WHERE v.modification_id = ANY(:ids) AND mv.month >= :start AND mv.month < :end GROUP BY 1, 2"
        ),
        {"ids": list(by_id), "start": start, "end": edge},
    ).all()

    table: dict[int, dict[float, dict[str, int]]] = defaultdict(lambda: defaultdict(lambda: defaultdict(int)))
    undated = 0

    def add(mod_id: int, month: date, score: float | None, count: int) -> None:
        nonlocal undated
        if score is None or count <= 0:
            return
        if requirement is not None and requirement_year(score, limits) != requirement:
            return
        if mod_id in undated_mods:
            undated += count
        key = f"{month.year}-{(month.month - 1) // 3 + 1}"
        table[by_id[mod_id]["manufacturer_id"]][float(score)][key] += count

    for mod_id, month, count in monthly:
        month_end = _add_months(month, 1)
        inside = [issued for issued, _ in segments[mod_id] if month < issued < month_end]
        if not inside:
            score = active(mod_id, month)
            if score is not False:
                add(mod_id, month, score, int(count))
            continue
        # Запись сменилась внутри месяца — делим поверки месяца по дням.
        days = db.execute(
            text(
                "SELECT mv.verification_date, count(*) FROM meter_verifications mv "
                "JOIN meter_market_variants v ON v.id = mv.variant_id WHERE v.modification_id = :m "
                "AND mv.applicable AND NOT mv.annulled AND mv.verification_date >= :s AND mv.verification_date < :e "
                "GROUP BY 1"
            ),
            {"m": mod_id, "s": month, "e": month_end},
        ).all()
        for day, day_count in days:
            score = active(mod_id, day)
            if score is not False:
                add(mod_id, month, score, int(day_count))

    rows = []
    for manufacturer_id, scores in table.items():
        by_q = {q: sum(s.get(q, 0) for s in scores.values()) for q in quarters}
        rows.append(
            {
                "manufacturer_id": manufacturer_id,
                "name": names.get(manufacturer_id, "Производитель не указан"),
                "quarters": by_q,
                "total": sum(by_q.values()),
                "scores": [
                    {
                        "score": score,
                        "requirement_year": requirement_year(score, limits),
                        "quarters": {q: values.get(q, 0) for q in quarters},
                        "total": sum(values.values()),
                    }
                    for score, values in sorted(scores.items())
                ],
            }
        )
    rows.sort(key=lambda r: -r["total"])
    totals = {q: sum(r["quarters"][q] for r in rows) for q in quarters}
    return {**empty, "rows": rows, "totals": totals, "total": sum(totals.values()), "undated": undated}


# ── Переход вглубь: производитель → тип → модификация → варианты ───────────────────────────


def drill(
    db: Session,
    period: Period,
    *,
    manufacturer_id: int | None = None,
    type_id: int | None = None,
    modification_id: int | None = None,
    unmapped_type_id: int | None = None,
) -> dict[str, Any]:
    params: dict[str, Any] = {"start": period.start, "end": period.end}
    period_where = "mv.month >= :start AND mv.month < :end"
    if modification_id is not None or unmapped_type_id is not None:
        if modification_id is not None:
            where, params["id"] = "v.modification_id = :id", modification_id
        else:
            where, params["id"] = "v.type_id = :id AND v.modification_id IS NULL", unmapped_type_id
        rows = db.execute(
            text(
                "SELECT v.id, v.raw, v.mapped_by, coalesce(sum(mv.count) FILTER (WHERE " + period_where + "), 0) AS count "
                "FROM meter_market_variants v LEFT JOIN meter_verification_monthly mv ON mv.variant_id = v.id "
                f"WHERE {where} GROUP BY 1, 2, 3 ORDER BY 4 DESC"
            ),
            params,
        ).mappings()
        return {"level": "variants", "rows": [dict(r) for r in rows]}
    if type_id is not None:
        rows = db.execute(
            text(
                "SELECT m.id, coalesce(m.name, :unmapped) AS name, m.phase, m.design, m.interface, sum(mv.count) AS count "
                f"{_BASE_FROM} WHERE t.id = :t AND {period_where} GROUP BY 1, 2, 3, 4, 5 ORDER BY 6 DESC"
            ),
            {**params, "t": type_id, "unmapped": UNMAPPED},
        ).mappings()
        return {"level": "modifications", "type_id": type_id, "rows": [dict(r) for r in rows]}
    if manufacturer_id is not None:
        rows = db.execute(
            text(
                "SELECT t.id, t.mit_number, t.notation, t.title, sum(mv.count) AS count "
                f"{_BASE_FROM} WHERE {_MANUFACTURER_SQL} = :m AND {period_where} GROUP BY 1, 2, 3, 4 ORDER BY 5 DESC"
            ),
            {**params, "m": manufacturer_id},
        ).mappings()
        return {"level": "types", "rows": [dict(r) for r in rows]}
    raise ValueError("Не указано, что раскрывать")


def variant_samples(db: Session, variant_id: int, period: Period, limit: int = 10) -> list[dict[str, Any]]:
    """Примеры поверок варианта — чтобы любое расхождение можно было показать записями ФГИС."""

    rows = db.execute(
        text(
            "SELECT mv.global_id, o.title AS org, mv.mi_number, mv.verification_date, mv.applicable, mv.annulled "
            "FROM meter_verifications mv LEFT JOIN meter_market_orgs o ON o.id = mv.org_id "
            "WHERE mv.variant_id = :v AND mv.verification_date >= :s AND mv.verification_date < :e "
            "ORDER BY mv.verification_date DESC LIMIT :l"
        ),
        {"v": variant_id, "s": period.start, "e": period.end, "l": limit},
    ).mappings()
    return [dict(r) for r in rows]


# ── Выгрузка в Excel ───────────────────────────────────────────────────────────────────────


def to_xlsx(title: str, subtitle: str, header: list[str], rows: list[list[Any]]) -> bytes:
    """Таблица экрана в Excel: шапка с названием и источником, автофильтр, закрепление шапки,
    числа — числами (их пересчитывают и сортируют, как и выгрузку тендеров, раздел 16)."""

    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = re.sub(r"[\\/*?:\[\]]", "", title)[:31] or "Отчёт"
    sheet.append([title])
    sheet["A1"].font = Font(bold=True, size=13)
    sheet.append([subtitle])
    sheet.append([])
    sheet.append(header)
    for cell in sheet[4]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    sheet.auto_filter.ref = f"A4:{get_column_letter(len(header))}{sheet.max_row}"
    sheet.freeze_panes = "A5"
    for index, name in enumerate(header, start=1):
        width = max([len(str(name))] + [len(str(r[index - 1])) for r in rows if index - 1 < len(r)])
        sheet.column_dimensions[get_column_letter(index)].width = min(60, max(10, width + 2))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
